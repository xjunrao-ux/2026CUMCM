"""A problem, question 4: moving-radius heat-moisture finite-volume model.

The cylinder radius is read from Attachment 2 and linearly interpolated in
time.  Uniform radial shrinkage is assumed, so xi=r/R(t) labels a material
position.  The material derivative therefore becomes a time derivative at
fixed xi, while the current radius remains in the physical finite-volume
geometry.  Appendix 4 properties are used throughout.  Temperature is stored
and advanced in kelvin, including the Arrhenius diffusivity evaluation.

The requested result4 grid is reported at fixed physical distances
0.0, 0.1, ..., 2.0 cm plus a separate moving-surface column.  A fixed-distance
cell is left blank after the shrinking surface has moved inside that distance.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

import a2_coupled_fvm as q2


INITIAL_RADIUS_M = 0.02
INITIAL_TEMPERATURE_K = 28.0 + 273.15
INITIAL_MOISTURE_KG_KG = 2.55
HEAT_TRANSFER_COEFFICIENT_W_M2_K = 25.0
MASS_TRANSFER_COEFFICIENT_M_S = 8.0e-7

MEASURED_ENVIRONMENT_END_S = 4.0 * 3600.0
PLATEAU_START_S = 3.0 * 3600.0
DRYING_THRESHOLD_KG_KG = 0.15
OUTPUT_INTERVAL_S = 60.0
TABLE_INTERVAL_S = 6.0 * 3600.0
MAXIMUM_SIMULATION_TIME_S = 7.0 * 24.0 * 3600.0

OUTPUT_RADII_CM = np.round(np.arange(0.0, 2.0 + 0.05, 0.1), 1)
TABLE_RADII_CM = np.round(np.arange(0.0, 2.0 + 0.25, 0.5), 1)


@dataclass(frozen=True)
class BoundaryProgram:
    source_times_s: np.ndarray
    source_temperature_k: np.ndarray
    source_moisture_kg_kg: np.ndarray
    plateau_temperature_k: float
    plateau_moisture_kg_kg: float

    def values(self, time_s: float) -> tuple[float, float]:
        if time_s <= MEASURED_ENVIRONMENT_END_S:
            return (
                float(
                    np.interp(time_s, self.source_times_s, self.source_temperature_k)
                ),
                float(
                    np.interp(
                        time_s,
                        self.source_times_s,
                        self.source_moisture_kg_kg,
                    )
                ),
            )
        return self.plateau_temperature_k, self.plateau_moisture_kg_kg


@dataclass(frozen=True)
class RadiusProgram:
    source_times_s: np.ndarray
    source_radii_cm: np.ndarray

    def radius_cm(self, time_s: float) -> float:
        return float(
            np.interp(
                time_s,
                self.source_times_s,
                self.source_radii_cm,
                left=self.source_radii_cm[0],
                right=self.source_radii_cm[-1],
            )
        )

    def radius_m(self, time_s: float) -> float:
        return self.radius_cm(time_s) / 100.0


@dataclass(frozen=True)
class MaterialGrid:
    centers_xi: np.ndarray
    west_faces_xi: np.ndarray
    east_faces_xi: np.ndarray


@dataclass
class StepResult:
    temperature_k: np.ndarray
    moisture_kg_kg: np.ndarray
    surface_temperature_k: float
    surface_moisture_kg_kg: float
    moisture_boundary_conductance: float
    moisture_storage_change: float
    minimum_diffusivity_m2_s: float
    maximum_diffusivity_m2_s: float
    iterations: int


@dataclass
class ShrinkageResult:
    times_s: np.ndarray
    radii_cm: np.ndarray
    temperature_k: np.ndarray
    moisture_kg_kg: np.ndarray
    surface_temperature_k: np.ndarray
    surface_moisture_kg_kg: np.ndarray
    drying_time_s: float
    previous_time_s: float
    previous_maximum_moisture: float
    final_maximum_moisture: float
    controlling_radius_cm: float
    final_internal_temperature_k: np.ndarray
    final_internal_moisture_kg_kg: np.ndarray
    maximum_picard_iterations: int
    time_step_counts: dict[str, int]
    discrete_moisture_storage_change: float
    integrated_moisture_outflow: float
    minimum_diffusivity_m2_s: float
    maximum_diffusivity_m2_s: float


def find_attachment(repo_root: Path, filename: str) -> Path:
    path = (
        repo_root
        / "比赛题目"
        / "CUMCM2026Problems"
        / "A题"
        / "附件"
        / filename
    )
    if not path.exists():
        raise FileNotFoundError(f"Cannot find {filename}: {path}")
    return path


def load_radius_program(path: Path) -> RadiusProgram:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [
        row[:2]
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    workbook.close()
    times = np.asarray([float(row[0]) for row in rows], dtype=float)
    radii = np.asarray([float(row[1]) for row in rows], dtype=float)
    if times.size < 2 or not np.all(np.diff(times) > 0.0):
        raise ValueError("Attachment 2 time values must be strictly increasing.")
    if not math.isclose(times[0], 0.0, abs_tol=1.0e-12):
        raise ValueError("Attachment 2 must start at t=0 s.")
    if not math.isclose(radii[0] / 100.0, INITIAL_RADIUS_M, abs_tol=1.0e-12):
        raise ValueError("Attachment 2 initial radius must be 2 cm.")
    if np.any(radii <= 0.0) or np.any(np.diff(radii) > 1.0e-12):
        raise ValueError("Attachment 2 radii must remain positive and non-increasing.")
    if not np.all(np.isfinite(radii)):
        raise ValueError("Attachment 2 contains a non-finite radius.")
    return RadiusProgram(times, radii)


def build_boundary_program(environment: q2.EnvironmentData) -> BoundaryProgram:
    """Use Attachment 1 through 4 h and its time-weighted 3-4 h mean later."""
    mask = (
        (environment.source_times_s >= PLATEAU_START_S)
        & (environment.source_times_s <= MEASURED_ENVIRONMENT_END_S)
    )
    times = environment.source_times_s[mask]
    temperature_k = environment.source_temperature_c[mask] + 273.15
    moisture = environment.source_moisture_kg_kg[mask]
    if (
        times.size < 2
        or not math.isclose(times[0], PLATEAU_START_S)
        or not math.isclose(times[-1], MEASURED_ENVIRONMENT_END_S)
    ):
        raise ValueError("Attachment 1 must contain the complete 3-4 h interval.")
    duration = MEASURED_ENVIRONMENT_END_S - PLATEAU_START_S
    return BoundaryProgram(
        source_times_s=environment.source_times_s,
        source_temperature_k=environment.source_temperature_c + 273.15,
        source_moisture_kg_kg=environment.source_moisture_kg_kg,
        plateau_temperature_k=float(np.trapezoid(temperature_k, times) / duration),
        plateau_moisture_kg_kg=float(np.trapezoid(moisture, times) / duration),
    )


def density_kg_m3(moisture: np.ndarray) -> np.ndarray:
    """Appendix 4 bulk density."""
    return 760.0 + 90.0 * np.asarray(moisture, dtype=float)


def heat_capacity_j_kg_k(moisture: np.ndarray) -> np.ndarray:
    """Appendix 4 specific heat capacity."""
    moisture = np.asarray(moisture, dtype=float)
    return 1850.0 + 2150.0 * moisture / (moisture + 1.0)


def conductivity_w_m_k(moisture: np.ndarray) -> np.ndarray:
    """Appendix 4 thermal conductivity."""
    moisture = np.asarray(moisture, dtype=float)
    return 0.12 + 0.20 * moisture / (moisture + 1.0)


def moisture_diffusivity_m2_s(
    moisture: np.ndarray, temperature_k: np.ndarray
) -> np.ndarray:
    """Appendix 4: D=4.2e-4 exp(-0.30/C) exp(-3850/T), with T in K."""
    moisture = np.asarray(moisture, dtype=float)
    temperature_k = np.asarray(temperature_k, dtype=float)
    if np.any(moisture <= 0.0) or np.any(temperature_k <= 0.0):
        raise ValueError("Positive moisture and absolute temperature are required.")
    return 4.2e-4 * np.exp(-0.30 / moisture) * np.exp(-3850.0 / temperature_k)


def make_material_grid(nominal_initial_step_cm: float) -> MaterialGrid:
    step_m = nominal_initial_step_cm / 100.0
    intervals = int(round(INITIAL_RADIUS_M / step_m))
    if not math.isclose(intervals * step_m, INITIAL_RADIUS_M, abs_tol=1.0e-12):
        raise ValueError("The nominal initial radial step must divide 2 cm exactly.")
    logical_faces = np.linspace(0.0, 1.0, intervals + 1)
    faces_xi = 1.0 - (1.0 - logical_faces) ** 2
    return MaterialGrid(
        centers_xi=0.5 * (faces_xi[:-1] + faces_xi[1:]),
        west_faces_xi=faces_xi[:-1],
        east_faces_xi=faces_xi[1:],
    )


def physical_grid(material_grid: MaterialGrid, radius_m: float) -> q2.Grid:
    west = radius_m * material_grid.west_faces_xi
    east = radius_m * material_grid.east_faces_xi
    centers = radius_m * material_grid.centers_xi
    volumes = math.pi * (east**2 - west**2)
    return q2.Grid(centers, west, east, volumes)


def build_dynamic_diffusion_system(
    grid: q2.Grid,
    radius_m: float,
    field_previous: np.ndarray,
    diffusion_coefficient: np.ndarray,
    storage_coefficient: np.ndarray,
    boundary_transfer_coefficient: float,
    boundary_value: float,
    time_step_s: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, float]:
    """Backward-Euler FVM system on the current shrunken physical geometry."""
    conductances = q2.internal_conductances(grid, diffusion_coefficient)
    storage = storage_coefficient * grid.volumes_m3_m / time_step_s
    lower = -conductances.copy()
    upper = -conductances.copy()
    diagonal = storage.copy()
    diagonal[:-1] += conductances
    diagonal[1:] += conductances
    rhs = storage * field_previous

    surface_area_per_length = 2.0 * math.pi * radius_m
    boundary_conductance = surface_area_per_length / (
        (radius_m - grid.radii_m[-1]) / diffusion_coefficient[-1]
        + 1.0 / boundary_transfer_coefficient
    )
    diagonal[-1] += boundary_conductance
    rhs[-1] += boundary_conductance * boundary_value
    return lower, diagonal, upper, rhs, boundary_conductance


def reconstruct_center(grid: q2.Grid, values: np.ndarray) -> float:
    """Second-order reconstruction using even symmetry at r=0."""
    r0_squared = grid.radii_m[0] ** 2
    r1_squared = grid.radii_m[1] ** 2
    return float(
        (values[0] * r1_squared - values[1] * r0_squared)
        / (r1_squared - r0_squared)
    )


def advance_coupled_step(
    material_grid: MaterialGrid,
    radius_m: float,
    temperature_previous_k: np.ndarray,
    moisture_previous: np.ndarray,
    oven_temperature_k: float,
    oven_moisture: float,
    time_step_s: float,
    temperature_tolerance_k: float = 1.0e-8,
    moisture_tolerance: float = 1.0e-10,
    maximum_iterations: int = 30,
) -> StepResult:
    """Advance one material-coordinate time step by Picard and Thomas solves."""
    grid = physical_grid(material_grid, radius_m)
    temperature_iterate = temperature_previous_k.copy()
    moisture_iterate = moisture_previous.copy()
    for iteration in range(1, maximum_iterations + 1):
        conductivity = conductivity_w_m_k(moisture_iterate)
        heat_storage = density_kg_m3(moisture_iterate) * heat_capacity_j_kg_k(
            moisture_iterate
        )
        heat_system = build_dynamic_diffusion_system(
            grid,
            radius_m,
            temperature_previous_k,
            conductivity,
            heat_storage,
            HEAT_TRANSFER_COEFFICIENT_W_M2_K,
            oven_temperature_k,
            time_step_s,
        )
        temperature_updated = q2.solve_tridiagonal(*heat_system[:4])

        diffusivity = moisture_diffusivity_m2_s(
            moisture_iterate, temperature_updated
        )
        moisture_system = build_dynamic_diffusion_system(
            grid,
            radius_m,
            moisture_previous,
            diffusivity,
            np.ones_like(moisture_iterate),
            MASS_TRANSFER_COEFFICIENT_M_S,
            oven_moisture,
            time_step_s,
        )
        moisture_updated = q2.solve_tridiagonal(*moisture_system[:4])

        temperature_error = float(
            np.max(np.abs(temperature_updated - temperature_iterate))
        )
        moisture_error = float(np.max(np.abs(moisture_updated - moisture_iterate)))
        temperature_iterate = temperature_updated
        moisture_iterate = moisture_updated
        if (
            temperature_error < temperature_tolerance_k
            and moisture_error < moisture_tolerance
        ):
            break
    else:
        raise RuntimeError("Coupled Picard iteration did not converge.")

    final_conductivity = conductivity_w_m_k(moisture_iterate)
    final_diffusivity = moisture_diffusivity_m2_s(
        moisture_iterate, temperature_iterate
    )
    surface_distance = radius_m - grid.radii_m[-1]
    surface_temperature = q2.reconstruct_surface_value(
        temperature_iterate[-1],
        oven_temperature_k,
        final_conductivity[-1],
        HEAT_TRANSFER_COEFFICIENT_W_M2_K,
        surface_distance,
    )
    surface_moisture = q2.reconstruct_surface_value(
        moisture_iterate[-1],
        oven_moisture,
        final_diffusivity[-1],
        MASS_TRANSFER_COEFFICIENT_M_S,
        surface_distance,
    )
    final_moisture_system = build_dynamic_diffusion_system(
        grid,
        radius_m,
        moisture_previous,
        final_diffusivity,
        np.ones_like(moisture_iterate),
        MASS_TRANSFER_COEFFICIENT_M_S,
        oven_moisture,
        time_step_s,
    )
    boundary_conductance = final_moisture_system[-1]
    storage_change = float(
        np.sum(grid.volumes_m3_m * (moisture_iterate - moisture_previous))
    )
    return StepResult(
        temperature_k=temperature_iterate,
        moisture_kg_kg=moisture_iterate,
        surface_temperature_k=float(surface_temperature),
        surface_moisture_kg_kg=float(surface_moisture),
        moisture_boundary_conductance=float(boundary_conductance),
        moisture_storage_change=storage_change,
        minimum_diffusivity_m2_s=float(np.min(final_diffusivity)),
        maximum_diffusivity_m2_s=float(np.max(final_diffusivity)),
        iterations=iteration,
    )


def sample_physical_field(
    material_grid: MaterialGrid,
    radius_m: float,
    values: np.ndarray,
    surface_value: float,
    query_radii_cm: np.ndarray,
) -> np.ndarray:
    """Sample current physical radii; return NaN where a point is outside."""
    grid = physical_grid(material_grid, radius_m)
    extended_radii = np.concatenate(([0.0], grid.radii_m, [radius_m]))
    extended_values = np.concatenate(
        ([reconstruct_center(grid, values)], values, [surface_value])
    )
    query_m = np.asarray(query_radii_cm, dtype=float) / 100.0
    sampled = np.full(query_m.shape, np.nan, dtype=float)
    valid = query_m <= radius_m + 1.0e-12
    sampled[valid] = np.interp(
        np.minimum(query_m[valid], radius_m), extended_radii, extended_values
    )
    return sampled


def maximum_with_location(
    material_grid: MaterialGrid,
    radius_m: float,
    moisture: np.ndarray,
    surface_moisture: float,
) -> tuple[float, float]:
    grid = physical_grid(material_grid, radius_m)
    values = np.concatenate(
        ([reconstruct_center(grid, moisture)], moisture, [surface_moisture])
    )
    radii_cm = np.concatenate(
        ([0.0], grid.radii_m * 100.0, [radius_m * 100.0])
    )
    index = int(np.argmax(values))
    return float(values[index]), float(radii_cm[index])


def choose_time_step(
    time_s: float,
    maximum_moisture: float,
    early_step_s: float,
    late_step_s: float,
    final_step_s: float,
) -> tuple[float, str]:
    if time_s < MEASURED_ENVIRONMENT_END_S:
        return early_step_s, "measured_environment"
    if maximum_moisture < DRYING_THRESHOLD_KG_KG + 0.002:
        return final_step_s, "threshold_refinement"
    return late_step_s, "constant_environment"


def simulate_until_dry(
    boundary: BoundaryProgram,
    radius_program: RadiusProgram,
    nominal_initial_step_cm: float,
    early_step_s: float,
    late_step_s: float,
    final_step_s: float,
) -> ShrinkageResult:
    material_grid = make_material_grid(nominal_initial_step_cm)
    n = material_grid.centers_xi.size
    temperature = np.full(n, INITIAL_TEMPERATURE_K, dtype=float)
    moisture = np.full(n, INITIAL_MOISTURE_KG_KG, dtype=float)
    surface_temperature = INITIAL_TEMPERATURE_K
    surface_moisture = INITIAL_MOISTURE_KG_KG
    time_s = 0.0
    next_output_s = OUTPUT_INTERVAL_S

    record_times: list[float] = []
    record_radii: list[float] = []
    temperature_records: list[np.ndarray] = []
    moisture_records: list[np.ndarray] = []
    surface_temperature_records: list[float] = []
    surface_moisture_records: list[float] = []

    maximum_moisture = INITIAL_MOISTURE_KG_KG
    controlling_radius_cm = 0.0
    previous_time_s = 0.0
    previous_maximum = maximum_moisture
    maximum_picard_iterations = 0
    time_step_counts = {
        "measured_environment": 0,
        "constant_environment": 0,
        "threshold_refinement": 0,
        "crossing_recomputed": 0,
    }
    discrete_storage_change = 0.0
    integrated_outflow = 0.0
    minimum_diffusivity = math.inf
    maximum_diffusivity = 0.0
    forced_final_resolution = False

    while time_s < MAXIMUM_SIMULATION_TIME_S:
        if forced_final_resolution:
            proposed_step = final_step_s
            step_label = "threshold_refinement"
        else:
            proposed_step, step_label = choose_time_step(
                time_s,
                maximum_moisture,
                early_step_s,
                late_step_s,
                final_step_s,
            )
        time_step_s = proposed_step
        if (
            time_s < MEASURED_ENVIRONMENT_END_S
            < time_s + time_step_s
        ):
            time_step_s = MEASURED_ENVIRONMENT_END_S - time_s
        if time_s < next_output_s < time_s + time_step_s:
            time_step_s = next_output_s - time_s
        if time_s + time_step_s > MAXIMUM_SIMULATION_TIME_S:
            time_step_s = MAXIMUM_SIMULATION_TIME_S - time_s
        if time_step_s <= 0.0:
            raise RuntimeError("Non-positive adaptive time step encountered.")

        new_time_s = time_s + time_step_s
        new_radius_m = radius_program.radius_m(new_time_s)
        oven_temperature_k, oven_moisture = boundary.values(new_time_s)
        step = advance_coupled_step(
            material_grid,
            new_radius_m,
            temperature,
            moisture,
            oven_temperature_k,
            oven_moisture,
            time_step_s,
        )
        new_maximum, new_controlling_radius = maximum_with_location(
            material_grid,
            new_radius_m,
            step.moisture_kg_kg,
            step.surface_moisture_kg_kg,
        )
        if new_maximum < DRYING_THRESHOLD_KG_KG and time_step_s > final_step_s:
            forced_final_resolution = True
            time_step_counts["crossing_recomputed"] += 1
            continue

        discrete_storage_change += step.moisture_storage_change
        integrated_outflow += (
            time_step_s
            * step.moisture_boundary_conductance
            * (step.moisture_kg_kg[-1] - oven_moisture)
        )
        temperature = step.temperature_k
        moisture = step.moisture_kg_kg
        surface_temperature = step.surface_temperature_k
        surface_moisture = step.surface_moisture_kg_kg
        time_s = new_time_s
        maximum_moisture = new_maximum
        controlling_radius_cm = new_controlling_radius
        maximum_picard_iterations = max(
            maximum_picard_iterations, step.iterations
        )
        minimum_diffusivity = min(minimum_diffusivity, step.minimum_diffusivity_m2_s)
        maximum_diffusivity = max(maximum_diffusivity, step.maximum_diffusivity_m2_s)
        time_step_counts[step_label] += 1

        is_output_time = math.isclose(time_s, next_output_s, abs_tol=1.0e-8)
        if is_output_time:
            current_radius_cm = new_radius_m * 100.0
            record_times.append(time_s)
            record_radii.append(current_radius_cm)
            temperature_records.append(
                sample_physical_field(
                    material_grid,
                    new_radius_m,
                    temperature,
                    surface_temperature,
                    OUTPUT_RADII_CM,
                )
            )
            moisture_records.append(
                sample_physical_field(
                    material_grid,
                    new_radius_m,
                    moisture,
                    surface_moisture,
                    OUTPUT_RADII_CM,
                )
            )
            surface_temperature_records.append(surface_temperature)
            surface_moisture_records.append(surface_moisture)
            next_output_s += OUTPUT_INTERVAL_S

        if maximum_moisture < DRYING_THRESHOLD_KG_KG:
            if not is_output_time:
                current_radius_cm = new_radius_m * 100.0
                record_times.append(time_s)
                record_radii.append(current_radius_cm)
                temperature_records.append(
                    sample_physical_field(
                        material_grid,
                        new_radius_m,
                        temperature,
                        surface_temperature,
                        OUTPUT_RADII_CM,
                    )
                )
                moisture_records.append(
                    sample_physical_field(
                        material_grid,
                        new_radius_m,
                        moisture,
                        surface_moisture,
                        OUTPUT_RADII_CM,
                    )
                )
                surface_temperature_records.append(surface_temperature)
                surface_moisture_records.append(surface_moisture)
            break

        previous_time_s = time_s
        previous_maximum = maximum_moisture
    else:
        raise RuntimeError("Maximum simulation time reached before drying.")

    return ShrinkageResult(
        times_s=np.asarray(record_times, dtype=float),
        radii_cm=np.asarray(record_radii, dtype=float),
        temperature_k=np.vstack(temperature_records),
        moisture_kg_kg=np.vstack(moisture_records),
        surface_temperature_k=np.asarray(surface_temperature_records, dtype=float),
        surface_moisture_kg_kg=np.asarray(surface_moisture_records, dtype=float),
        drying_time_s=time_s,
        previous_time_s=previous_time_s,
        previous_maximum_moisture=previous_maximum,
        final_maximum_moisture=maximum_moisture,
        controlling_radius_cm=controlling_radius_cm,
        final_internal_temperature_k=temperature,
        final_internal_moisture_kg_kg=moisture,
        maximum_picard_iterations=maximum_picard_iterations,
        time_step_counts=time_step_counts,
        discrete_moisture_storage_change=discrete_storage_change,
        integrated_moisture_outflow=integrated_outflow,
        minimum_diffusivity_m2_s=minimum_diffusivity,
        maximum_diffusivity_m2_s=maximum_diffusivity,
    )


def record_indices_at_times(
    result: ShrinkageResult, query_times_s: np.ndarray
) -> np.ndarray:
    lookup = {round(float(time), 8): index for index, time in enumerate(result.times_s)}
    indices: list[int] = []
    for time in query_times_s:
        key = round(float(time), 8)
        if key not in lookup:
            raise ValueError(f"Requested time {time:g} s was not recorded.")
        indices.append(lookup[key])
    return np.asarray(indices, dtype=int)


def build_table_times(drying_time_s: float) -> np.ndarray:
    regular = np.arange(
        TABLE_INTERVAL_S,
        drying_time_s - 1.0e-9,
        TABLE_INTERVAL_S,
        dtype=float,
    )
    if regular.size and math.isclose(regular[-1], drying_time_s, abs_tol=1.0e-8):
        return regular
    return np.append(regular, drying_time_s)


def write_full_csv(path: Path, result: ShrinkageResult) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "time_s",
                "radius_cm",
                *[f"r_{radius:.1f}_cm" for radius in OUTPUT_RADII_CM],
                "surface",
            ]
        )
        for time_s, radius_cm, values, surface in zip(
            result.times_s,
            result.radii_cm,
            result.moisture_kg_kg,
            result.surface_moisture_kg_kg,
        ):
            writer.writerow(
                [
                    f"{time_s:.0f}",
                    f"{radius_cm:.6f}",
                    *["" if np.isnan(value) else f"{value:.4f}" for value in values],
                    f"{surface:.4f}",
                ]
            )


def write_temperature_csv(path: Path, result: ShrinkageResult) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "time_s",
                "radius_cm",
                *[f"r_{radius:.1f}_cm" for radius in OUTPUT_RADII_CM],
                "surface_temperature_K",
            ]
        )
        for time_s, radius_cm, values, surface in zip(
            result.times_s,
            result.radii_cm,
            result.temperature_k,
            result.surface_temperature_k,
        ):
            writer.writerow(
                [
                    f"{time_s:.0f}",
                    f"{radius_cm:.6f}",
                    *["" if np.isnan(value) else f"{value:.4f}" for value in values],
                    f"{surface:.4f}",
                ]
            )


def write_table_csv(
    path: Path,
    result: ShrinkageResult,
    table_times_s: np.ndarray,
) -> None:
    indices = record_indices_at_times(result, table_times_s)
    output_columns = [
        int(np.where(np.isclose(OUTPUT_RADII_CM, radius))[0][0])
        for radius in TABLE_RADII_CM
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            ["时间/h", *[f"{radius:g} cm" for radius in TABLE_RADII_CM], "药材表面"]
        )
        for time_s, index in zip(table_times_s, indices):
            label = (
                "烘干结束时间"
                if math.isclose(time_s, result.drying_time_s, abs_tol=1.0e-8)
                else f"{time_s / 3600.0:.0f}"
            )
            values = result.moisture_kg_kg[index, output_columns]
            writer.writerow(
                [
                    label,
                    *["" if np.isnan(value) else f"{value:.4f}" for value in values],
                    f"{result.surface_moisture_kg_kg[index]:.4f}",
                ]
            )


def write_workbook_payload(path: Path, result: ShrinkageResult) -> None:
    rows: list[list[float | None]] = []
    for time_s, values, surface in zip(
        result.times_s, result.moisture_kg_kg, result.surface_moisture_kg_kg
    ):
        rows.append(
            [
                int(round(float(time_s))),
                *[None if np.isnan(value) else round(float(value), 4) for value in values],
                round(float(surface), 4),
            ]
        )
    payload = {
        "headers": [
            "时间\\到药材中心的距离",
            *[round(float(radius), 1) for radius in OUTPUT_RADII_CM],
            "药材表面",
        ],
        "rows": rows,
        "drying_time_s": result.drying_time_s,
    }
    with path.open("w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, separators=(",", ":"))


def common_refinement_difference(
    production: ShrinkageResult, coarse: ShrinkageResult
) -> tuple[int, float, float]:
    coarse_lookup = {
        round(float(time), 8): index for index, time in enumerate(coarse.times_s)
    }
    differences: list[np.ndarray] = []
    common_rows = 0
    for index, time in enumerate(production.times_s):
        key = round(float(time), 8)
        if key not in coarse_lookup:
            continue
        coarse_index = coarse_lookup[key]
        a = production.moisture_kg_kg[index]
        b = coarse.moisture_kg_kg[coarse_index]
        mask = np.isfinite(a) & np.isfinite(b)
        if np.any(mask):
            differences.append(np.abs(a[mask] - b[mask]))
            differences.append(
                np.asarray(
                    [
                        abs(
                            production.surface_moisture_kg_kg[index]
                            - coarse.surface_moisture_kg_kg[coarse_index]
                        )
                    ]
                )
            )
            common_rows += 1
    combined = np.concatenate(differences)
    return common_rows, float(np.max(combined)), float(np.mean(combined))


def build_validation(
    boundary: BoundaryProgram,
    radius_program: RadiusProgram,
    production: ShrinkageResult,
    coarse: ShrinkageResult,
) -> dict:
    common_rows, refinement_maximum, refinement_mean = common_refinement_difference(
        production, coarse
    )
    residual = (
        production.discrete_moisture_storage_change
        + production.integrated_moisture_outflow
    )
    balance_scale = max(
        abs(production.discrete_moisture_storage_change),
        abs(production.integrated_moisture_outflow),
        1.0e-30,
    )

    invalid_distance_values = 0
    missing_inside_values = 0
    radial_order_violations = 0
    time_order_violations = 0
    for row_index, radius_cm in enumerate(production.radii_cm):
        inside = OUTPUT_RADII_CM <= radius_cm + 1.0e-10
        invalid_distance_values += int(
            np.count_nonzero(np.isfinite(production.moisture_kg_kg[row_index, ~inside]))
        )
        missing_inside_values += int(
            np.count_nonzero(~np.isfinite(production.moisture_kg_kg[row_index, inside]))
        )
        radial_values = np.concatenate(
            (
                production.moisture_kg_kg[row_index, inside],
                [production.surface_moisture_kg_kg[row_index]],
            )
        )
        radial_order_violations += int(np.count_nonzero(np.diff(radial_values) > 1.0e-8))
    for column in range(OUTPUT_RADII_CM.size):
        series = production.moisture_kg_kg[:, column]
        valid = np.isfinite(series)
        if np.count_nonzero(valid) > 1:
            time_order_violations += int(
                np.count_nonzero(np.diff(series[valid]) > 1.0e-8)
            )
    time_order_violations += int(
        np.count_nonzero(np.diff(production.surface_moisture_kg_kg) > 1.0e-8)
    )

    source_reconstruction_error = float(
        np.max(
            np.abs(
                np.asarray(
                    [radius_program.radius_cm(t) for t in radius_program.source_times_s]
                )
                - radius_program.source_radii_cm
            )
        )
    )
    finite_temperature = production.temperature_k[np.isfinite(production.temperature_k)]
    finite_moisture = production.moisture_kg_kg[np.isfinite(production.moisture_kg_kg)]
    return {
        "model_scope": {
            "geometry": "one-dimensional axisymmetric cylinder with moving radius",
            "coordinate": "uniform material coordinate xi=r/R(t)",
            "radius_interpolation": "piecewise linear Attachment 2 data; last value held only if needed",
            "constitutive_laws": "Appendix 4 used over the full process",
            "temperature_state_unit": "K",
            "latent_heat": "excluded consistently with the simpler questions 1-3 model",
            "drying_threshold_kg_kg": DRYING_THRESHOLD_KG_KG,
            "environment_after_4h": "time-weighted mean over Attachment 1 final hour",
            "result4_distance_rule": (
                "fixed physical radii 0.0-2.0 cm every 0.1 cm plus the moving "
                "surface; a fixed-radius cell is blank when its radius exceeds R(t)"
            ),
            "production_initial_nominal_step_cm": 0.00625,
            "production_time_steps_s": [1.0, 30.0, 1.0],
            "coarse_initial_nominal_step_cm": 0.0125,
            "coarse_time_steps_s": [1.0, 30.0, 1.0],
            "saved_interval_s": OUTPUT_INTERVAL_S,
        },
        "constant_boundary": {
            "temperature_K": boundary.plateau_temperature_k,
            "temperature_C_for_reference": boundary.plateau_temperature_k - 273.15,
            "moisture_kg_kg": boundary.plateau_moisture_kg_kg,
        },
        "radius": {
            "attachment_interval_s": [
                float(radius_program.source_times_s[0]),
                float(radius_program.source_times_s[-1]),
            ],
            "initial_cm": float(radius_program.source_radii_cm[0]),
            "drying_time_cm": float(production.radii_cm[-1]),
            "source_node_reconstruction_max_error_cm": source_reconstruction_error,
            "recorded_nonincreasing": bool(np.all(np.diff(production.radii_cm) <= 1.0e-12)),
        },
        "drying_time": {
            "production_s": production.drying_time_s,
            "production_h": production.drying_time_s / 3600.0,
            "coarse_s": coarse.drying_time_s,
            "coarse_h": coarse.drying_time_s / 3600.0,
            "coarse_production_difference_s": abs(
                production.drying_time_s - coarse.drying_time_s
            ),
            "previous_time_s": production.previous_time_s,
            "previous_maximum_moisture": production.previous_maximum_moisture,
            "final_maximum_moisture": production.final_maximum_moisture,
            "controlling_radius_cm": production.controlling_radius_cm,
        },
        "grid_refinement": {
            "common_record_count": common_rows,
            "maximum_moisture_change_kg_kg": refinement_maximum,
            "mean_moisture_change_kg_kg": refinement_mean,
        },
        "discrete_moisture_balance": {
            "summed_storage_change": production.discrete_moisture_storage_change,
            "integrated_boundary_outflow": production.integrated_moisture_outflow,
            "residual": residual,
            "relative_residual": abs(residual) / balance_scale,
        },
        "physical_ranges": {
            "temperature_min_K": float(np.min(finite_temperature)),
            "temperature_max_K": float(np.max(finite_temperature)),
            "moisture_min_kg_kg": float(np.min(finite_moisture)),
            "moisture_max_kg_kg": float(np.max(finite_moisture)),
            "surface_moisture_min_kg_kg": float(np.min(production.surface_moisture_kg_kg)),
            "diffusivity_min_m2_s": production.minimum_diffusivity_m2_s,
            "diffusivity_max_m2_s": production.maximum_diffusivity_m2_s,
        },
        "checks": {
            "previous_state_not_dry": bool(
                production.previous_maximum_moisture >= DRYING_THRESHOLD_KG_KG
            ),
            "final_state_all_dry": bool(
                production.final_maximum_moisture < DRYING_THRESHOLD_KG_KG
            ),
            "invalid_values_outside_current_radius": invalid_distance_values,
            "missing_values_inside_current_radius": missing_inside_values,
            "radial_order_violation_count": radial_order_violations,
            "time_monotonicity_violation_count": time_order_violations,
            "maximum_picard_iterations": production.maximum_picard_iterations,
            "time_step_counts": production.time_step_counts,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment1", type=Path, help="Path to Attachment 1")
    parser.add_argument("--attachment2", type=Path, help="Path to Attachment 2")
    parser.add_argument("--repo-root", type=Path, help="Repository root")
    args = parser.parse_args()

    script_path = Path(__file__).resolve()
    repo_root = args.repo_root.resolve() if args.repo_root else script_path.parents[1]
    attachment1 = (
        args.attachment1.resolve()
        if args.attachment1
        else find_attachment(repo_root, "附件1.xlsx")
    )
    attachment2 = (
        args.attachment2.resolve()
        if args.attachment2
        else find_attachment(repo_root, "附件2.xlsx")
    )
    result_dir = repo_root / "results" / "A_problem4_shrinkage"
    result_dir.mkdir(parents=True, exist_ok=True)

    environment = q2.load_and_preprocess_environment(attachment1)
    boundary = build_boundary_program(environment)
    radius_program = load_radius_program(attachment2)

    coarse = simulate_until_dry(
        boundary,
        radius_program,
        nominal_initial_step_cm=0.0125,
        early_step_s=1.0,
        late_step_s=30.0,
        final_step_s=1.0,
    )
    production = simulate_until_dry(
        boundary,
        radius_program,
        nominal_initial_step_cm=0.00625,
        early_step_s=1.0,
        late_step_s=30.0,
        final_step_s=1.0,
    )

    table_times = build_table_times(production.drying_time_s)
    write_full_csv(result_dir / "moisture_full_60s_0p1cm.csv", production)
    write_temperature_csv(
        result_dir / "temperature_auxiliary_60s_0p1cm_K.csv", production
    )
    write_table_csv(result_dir / "table6_moisture.csv", production, table_times)
    write_workbook_payload(result_dir / "result4_payload.json", production)
    validation = build_validation(boundary, radius_program, production, coarse)
    with (result_dir / "validation_summary.json").open("w", encoding="utf-8") as stream:
        json.dump(validation, stream, ensure_ascii=False, indent=2)

    print(
        f"Drying time: {production.drying_time_s:.0f} s "
        f"= {production.drying_time_s / 3600.0:.6f} h"
    )
    print(
        f"Threshold bracket: {production.previous_time_s:.0f} -> "
        f"{production.drying_time_s:.0f} s; "
        f"{production.previous_maximum_moisture:.10f} -> "
        f"{production.final_maximum_moisture:.10f} kg/kg"
    )
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    print(f"Results: {result_dir}")


if __name__ == "__main__":
    main()
