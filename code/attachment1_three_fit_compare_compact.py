"""附件1在0--1800 s内的三种边界曲线构造方法对比。

比较分段线性插值、PCHIP和端点约束的平移拉伸指数模型。
拉伸指数模型：
    y(t)=y0+(yT-y0) g(t;tau,beta)
其中 g(0)=0、g(T)=1，tau>0、0.2<=beta<=1；使用60 s正时间平移
避免beta<1时经典模型在t=0出现无穷导数。
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass, asdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator
from scipy.optimize import least_squares


END_TIME_S = 1800.0
TIME_SHIFT_S = 60.0
BETA_BOUNDS = (0.2, 1.0)
TAU_BOUNDS_S = (60.0, 1.0e6)


@dataclass
class StretchFit:
    tau_s: float
    beta: float
    success: bool
    cost: float


def configure_fonts() -> None:
    available = {font.name for font in fm.fontManager.ttflist}
    candidates = ["Microsoft YaHei", "SimHei", "Arial", "DejaVu Sans"]
    selected = [name for name in candidates if name in available]
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": selected or ["DejaVu Sans"],
        "axes.unicode_minus": False,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "font.size": 8.0,
        "axes.labelsize": 8.0,
        "axes.titlesize": 9.0,
        "xtick.labelsize": 7.0,
        "ytick.labelsize": 7.0,
        "legend.fontsize": 7.0,
        "axes.linewidth": 0.7,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "legend.frameon": False,
    })


def load_attachment(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [
        row[:3] for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None and float(row[0]) <= END_TIME_S
    ]
    workbook.close()
    data = np.asarray(rows, dtype=float)
    times_s, temperature_c, moisture_kg_kg = data.T
    if times_s[0] != 0.0 or times_s[-1] != END_TIME_S:
        raise ValueError("附件1没有完整覆盖0--1800 s。")
    if not np.all(np.diff(times_s) > 0.0):
        raise ValueError("附件1时间必须严格递增。")
    return times_s, temperature_c, moisture_kg_kg


def stretch_shape(times_s: np.ndarray, tau_s: float, beta: float) -> np.ndarray:
    """Stable normalized shifted stretched-exponential shape on [0,T]."""
    times = np.asarray(times_s, dtype=float)
    a0 = (TIME_SHIFT_S / tau_s) ** beta
    delta = ((times + TIME_SHIFT_S) / tau_s) ** beta - a0
    delta_end = ((END_TIME_S + TIME_SHIFT_S) / tau_s) ** beta - a0
    denominator = -np.expm1(-delta_end)
    if not np.isfinite(denominator) or denominator <= 0.0:
        raise ValueError("拉伸指数归一化分母无效。")
    shape = -np.expm1(-delta) / denominator
    # Only pin entries that are actually endpoints; prediction arrays in
    # leave-one-out validation may contain a single interior time.
    shape[np.isclose(times, 0.0)] = 0.0
    shape[np.isclose(times, END_TIME_S)] = 1.0
    return shape


def evaluate_stretched(
    times_s: np.ndarray, y0: float, y_end: float, tau_s: float, beta: float
) -> np.ndarray:
    return y0 + (y_end - y0) * stretch_shape(times_s, tau_s, beta)


def fit_stretched(times_s: np.ndarray, values: np.ndarray) -> StretchFit:
    y0, y_end = float(values[0]), float(values[-1])
    scale = max(float(np.ptp(values)), 1.0e-12)
    if TAU_BOUNDS_S[0] <= 0.0 or TAU_BOUNDS_S[1] <= TAU_BOUNDS_S[0]:
        raise ValueError("tau的对数优化边界必须为正且严格递增。")

    def residual(parameters: np.ndarray) -> np.ndarray:
        tau_s = float(np.exp(parameters[0]))
        beta = float(parameters[1])
        fitted = evaluate_stretched(times_s, y0, y_end, tau_s, beta)
        return (fitted - values) / scale

    lower = np.array([np.log(TAU_BOUNDS_S[0]), BETA_BOUNDS[0]])
    upper = np.array([np.log(TAU_BOUNDS_S[1]), BETA_BOUNDS[1]])
    candidates = []
    for tau0 in [300.0, 900.0, 1800.0, 6000.0, 30000.0, 300000.0]:
        if tau0 <= 0.0:
            raise ValueError("tau初值必须为正。")
        for beta0 in [0.35, 0.60, 0.85, 0.98]:
            result = least_squares(
                residual, x0=np.array([np.log(tau0), beta0]),
                bounds=(lower, upper), xtol=1.0e-12, ftol=1.0e-12,
                gtol=1.0e-12, max_nfev=5000,
            )
            candidates.append(result)
    best = min(candidates, key=lambda item: 2.0 * item.cost)
    return StretchFit(
        tau_s=float(np.exp(best.x[0])), beta=float(best.x[1]),
        success=bool(best.success), cost=float(2.0 * best.cost),
    )


def fit_all(
    times_s: np.ndarray, values: np.ndarray, query_times_s: np.ndarray
) -> tuple[dict[str, np.ndarray], StretchFit]:
    fit = fit_stretched(times_s, values)
    predictions = {
        "linear": np.interp(query_times_s, times_s, values),
        "pchip": PchipInterpolator(times_s, values)(query_times_s),
        "stretched": evaluate_stretched(
            query_times_s, float(values[0]), float(values[-1]), fit.tau_s, fit.beta
        ),
    }
    return predictions, fit


def loocv_predictions(times_s: np.ndarray, values: np.ndarray) -> dict[str, np.ndarray]:
    """Leave out each interior knot; endpoints remain to preserve the interval."""
    predictions = {name: np.full(values.size, np.nan) for name in ["linear", "pchip", "stretched"]}
    for index in range(1, len(times_s) - 1):
        keep = np.ones(times_s.size, dtype=bool)
        keep[index] = False
        train_t, train_y = times_s[keep], values[keep]
        query = float(times_s[index])
        predictions["linear"][index] = float(np.interp(query, train_t, train_y))
        predictions["pchip"][index] = float(PchipInterpolator(train_t, train_y)(query))
        fit = fit_stretched(train_t, train_y)
        predictions["stretched"][index] = float(evaluate_stretched(
            np.asarray([query]), float(train_y[0]), float(train_y[-1]), fit.tau_s, fit.beta
        )[0])
    return predictions


def segment_overshoot_count(
    times_s: np.ndarray, observed: np.ndarray,
    query_times_s: np.ndarray, predicted: np.ndarray,
    require_interpolation: bool,
) -> int:
    """Count interpolation overshoots only for methods that claim knot interpolation."""
    if not require_interpolation:
        return 0
    segment = np.searchsorted(times_s, query_times_s, side="right") - 1
    segment = np.clip(segment, 0, len(times_s) - 2)
    lower = np.minimum(observed[segment], observed[segment + 1])
    upper = np.maximum(observed[segment], observed[segment + 1])
    return int(np.count_nonzero((predicted < lower - 1e-12) | (predicted > upper + 1e-12)))


def build_metrics(
    variable: str, unit: str, times_s: np.ndarray, observed: np.ndarray,
    query_times_s: np.ndarray, predictions: dict[str, np.ndarray],
    loocv: dict[str, np.ndarray],
) -> list[dict]:
    knot_index = np.searchsorted(query_times_s, times_s)
    rows = []
    for method, predicted in predictions.items():
        knot_error = predicted[knot_index] - observed
        cv_error = loocv[method][1:-1] - observed[1:-1]
        monotonic_violations = int(np.count_nonzero(np.diff(predicted) < -1.0e-12))
        if method == "linear":
            slopes = np.diff(observed) / np.diff(times_s)
            max_slope_jump = float(np.max(np.abs(np.diff(slopes))))
            slope_discontinuity_count = int(np.count_nonzero(np.abs(np.diff(slopes)) > 1.0e-12))
        else:
            max_slope_jump = 0.0
            slope_discontinuity_count = 0
        rows.append({
            "variable": variable,
            "unit": unit,
            "method": method,
            "observation_count": int(observed.size),
            "interior_loocv_count": int(observed.size - 2),
            "knot_mae": float(np.mean(np.abs(knot_error))),
            "knot_rmse": float(np.sqrt(np.mean(knot_error**2))),
            "knot_max_abs_error": float(np.max(np.abs(knot_error))),
            "loocv_mae": float(np.mean(np.abs(cv_error))),
            "loocv_rmse": float(np.sqrt(np.mean(cv_error**2))),
            "loocv_max_abs_error": float(np.max(np.abs(cv_error))),
            "endpoint_max_abs_error": float(max(abs(knot_error[0]), abs(knot_error[-1]))),
            "monotonicity_violation_count": monotonic_violations,
            "local_overshoot_count": segment_overshoot_count(
                times_s, observed, query_times_s, predicted,
                require_interpolation=method in {"linear", "pchip"},
            ),
            "slope_discontinuity_count": slope_discontinuity_count,
            "maximum_slope_jump_per_s": max_slope_jump,
        })
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_full_comparison(
    path: Path, source_times_s: np.ndarray, source_temperature: np.ndarray,
    source_moisture: np.ndarray, query_times_s: np.ndarray,
    temperature: dict[str, np.ndarray], moisture: dict[str, np.ndarray],
) -> None:
    lookup = {int(t): (float(a), float(b)) for t, a, b in zip(source_times_s, source_temperature, source_moisture)}
    rows = []
    for i, t in enumerate(query_times_s.astype(int)):
        observed = lookup.get(t)
        rows.append({
            "time_s": t,
            "is_observation": int(observed is not None),
            "observed_temperature_C": "" if observed is None else observed[0],
            "temperature_linear_C": temperature["linear"][i],
            "temperature_pchip_C": temperature["pchip"][i],
            "temperature_stretched_C": temperature["stretched"][i],
            "temperature_pchip_minus_linear_C": temperature["pchip"][i] - temperature["linear"][i],
            "temperature_stretched_minus_linear_C": temperature["stretched"][i] - temperature["linear"][i],
            "observed_moisture_kg_kg": "" if observed is None else observed[1],
            "moisture_linear_kg_kg": moisture["linear"][i],
            "moisture_pchip_kg_kg": moisture["pchip"][i],
            "moisture_stretched_kg_kg": moisture["stretched"][i],
            "moisture_pchip_minus_linear_kg_kg": moisture["pchip"][i] - moisture["linear"][i],
            "moisture_stretched_minus_linear_kg_kg": moisture["stretched"][i] - moisture["linear"][i],
        })
    write_csv(path, rows)


def style_axis(ax: plt.Axes) -> None:
    ax.grid(True, color="#D9D9D9", linewidth=0.45, alpha=0.8)
    ax.set_axisbelow(True)


def plot_comparison(
    path: Path, times_s: np.ndarray, temperature_obs: np.ndarray,
    moisture_obs: np.ndarray, query_times_s: np.ndarray,
    temperature: dict[str, np.ndarray], moisture: dict[str, np.ndarray],
) -> None:
    configure_fonts()
    fig, axes = plt.subplots(2, 2, figsize=(7.20, 4.75), constrained_layout=True,
                             gridspec_kw={"height_ratios": [1.25, 1.0]})
    minute = query_times_s / 60.0
    obs_minute = times_s / 60.0
    colors = {"linear": "#0072B2", "pchip": "#E69F00", "stretched": "#CC79A7"}
    labels = {"linear": "分段线性", "pchip": "PCHIP", "stretched": "约束拉伸指数"}
    styles = {"linear": "-", "pchip": "--", "stretched": "-."}

    for col, (observed, prediction, unit, title) in enumerate([
        (temperature_obs, temperature, "°C", "烘房温度"),
        (moisture_obs, moisture, "kg/kg", "烘房水分浓度"),
    ]):
        for method in ["linear", "pchip", "stretched"]:
            axes[0, col].plot(minute, prediction[method], color=colors[method],
                              ls=styles[method], lw=1.35, label=labels[method])
        axes[0, col].scatter(obs_minute, observed, s=9, color="#222222", zorder=4,
                             label="附件1观测点")
        axes[0, col].set(
            title=f"({'ab'[col]}) {title}拟合",
            ylabel=f"数值 / ({unit})",
            xlim=(0, 30),
        )
        axes[0, col].legend(loc="best", ncol=1)

        for method in ["pchip", "stretched"]:
            difference = prediction[method] - prediction["linear"]
            axes[1, col].plot(minute, difference, color=colors[method],
                              ls=styles[method], lw=1.1, label=f"{labels[method]}−线性")
        axes[1, col].axhline(0, color="#555555", lw=0.7)
        axes[1, col].set(
            title=f"({'cd'[col]}) 相对分段线性的差值",
            xlabel="时间 / min",
            ylabel=f"差值 / ({unit})",
            xlim=(0, 30),
        )
        axes[1, col].legend(loc="best")

    for ax in axes.ravel():
        style_axis(ax)
        ax.set_xticks(np.arange(0, 31, 5))
    fig.savefig(path, dpi=600, facecolor="white")
    plt.close(fig)


def plot_metrics_table(path: Path, metrics: list[dict], parameters: list[dict]) -> None:
    configure_fonts()
    fig, axes = plt.subplots(2, 1, figsize=(7.20, 4.30), constrained_layout=True,
                             gridspec_kw={"height_ratios": [2.0, 1.0]})
    for ax in axes:
        ax.axis("off")
    method_label = {"linear": "分段线性", "pchip": "PCHIP", "stretched": "约束拉伸指数"}
    rows = []
    for item in metrics:
        decimals = 6 if item["variable"] == "temperature" else 8
        rows.append([
            "温度" if item["variable"] == "temperature" else "水分浓度",
            method_label[item["method"]],
            f'{item["knot_rmse"]:.{decimals}f}',
            f'{item["loocv_rmse"]:.{decimals}f}',
            f'{item["loocv_mae"]:.{decimals}f}',
            str(item["monotonicity_violation_count"]),
            str(item["slope_discontinuity_count"]),
        ])
    headers = ["变量", "方法", "节点RMSE", "留一RMSE", "留一MAE", "单调违例", "斜率间断点"]
    table = axes[0].table(cellText=rows, colLabels=headers, cellLoc="center", loc="center",
                          colWidths=[0.13, 0.19, 0.15, 0.15, 0.15, 0.12, 0.14])
    table.auto_set_font_size(False)
    table.set_fontsize(6.8)
    table.scale(1.0, 1.42)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("#D0D0D0")
        cell.set_linewidth(0.45)
        if r == 0:
            cell.set_facecolor("#1F4E79")
            cell.get_text().set_color("white")
            cell.get_text().set_fontweight("bold")
        elif r % 2 == 0:
            cell.set_facecolor("#F3F6F8")
    axes[0].set_title("拟合精度、留一预测和形状约束", fontsize=9.4, fontweight="bold", pad=5)

    p_rows = [[
        "温度" if p["variable"] == "temperature" else "水分浓度",
        f'{p["tau_s"]:.2f}', f'{p["beta"]:.6f}',
        "是" if p["beta_at_bound"] else "否",
        "首末端点固定；单调；60 s时间平移",
    ] for p in parameters]
    p_headers = ["变量", "τ / s", "β", "β触及边界", "约束"]
    p_table = axes[1].table(cellText=p_rows, colLabels=p_headers, cellLoc="center", loc="center",
                            colWidths=[0.14, 0.16, 0.16, 0.16, 0.38])
    p_table.auto_set_font_size(False)
    p_table.set_fontsize(6.8)
    p_table.scale(1.0, 1.38)
    for (r, c), cell in p_table.get_celld().items():
        cell.set_edgecolor("#D0D0D0")
        cell.set_linewidth(0.45)
        if r == 0:
            cell.set_facecolor("#4F6D7A")
            cell.get_text().set_color("white")
            cell.get_text().set_fontweight("bold")
        elif r % 2 == 0:
            cell.set_facecolor("#F3F6F8")
    axes[1].set_title("约束拉伸指数参数", fontsize=9.0, fontweight="bold", pad=4)
    fig.savefig(path, dpi=600, facecolor="white")
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()

    result_dir = args.output_root / "results" / "A_problem1_interpolation"
    picture_dir = args.output_root / "picture" / "A_problem1_interpolation"
    result_dir.mkdir(parents=True, exist_ok=True)
    picture_dir.mkdir(parents=True, exist_ok=True)

    times_s, temperature_obs, moisture_obs = load_attachment(args.data)
    query_times_s = np.arange(0.0, END_TIME_S + 1.0)
    temperature, temp_fit = fit_all(times_s, temperature_obs, query_times_s)
    moisture, moisture_fit = fit_all(times_s, moisture_obs, query_times_s)
    temp_loocv = loocv_predictions(times_s, temperature_obs)
    moisture_loocv = loocv_predictions(times_s, moisture_obs)
    metrics = (
        build_metrics("temperature", "degC", times_s, temperature_obs, query_times_s, temperature, temp_loocv)
        + build_metrics("moisture", "kg/kg", times_s, moisture_obs, query_times_s, moisture, moisture_loocv)
    )
    parameters = [
        {"variable": "temperature", **asdict(temp_fit),
         "beta_at_bound": bool(abs(temp_fit.beta - BETA_BOUNDS[0]) < 1e-5 or abs(temp_fit.beta - BETA_BOUNDS[1]) < 1e-5),
         "time_shift_s": TIME_SHIFT_S, "beta_lower": BETA_BOUNDS[0], "beta_upper": BETA_BOUNDS[1],
         "tau_lower_s": TAU_BOUNDS_S[0], "tau_upper_s": TAU_BOUNDS_S[1]},
        {"variable": "moisture", **asdict(moisture_fit),
         "beta_at_bound": bool(abs(moisture_fit.beta - BETA_BOUNDS[0]) < 1e-5 or abs(moisture_fit.beta - BETA_BOUNDS[1]) < 1e-5),
         "time_shift_s": TIME_SHIFT_S, "beta_lower": BETA_BOUNDS[0], "beta_upper": BETA_BOUNDS[1],
         "tau_lower_s": TAU_BOUNDS_S[0], "tau_upper_s": TAU_BOUNDS_S[1]},
    ]

    write_full_comparison(result_dir / "attachment1_three_method_comparison_1s.csv",
                          times_s, temperature_obs, moisture_obs, query_times_s,
                          temperature, moisture)
    write_csv(result_dir / "attachment1_three_method_metrics.csv", metrics)
    write_csv(result_dir / "attachment1_stretched_exponential_parameters.csv", parameters)
    plot_comparison(picture_dir / "attachment1_three_method_comparison_compact.png",
                    times_s, temperature_obs, moisture_obs, query_times_s,
                    temperature, moisture)
    plot_metrics_table(picture_dir / "attachment1_three_method_metrics_table.png", metrics, parameters)

    summary = {"metrics": metrics, "parameters": parameters}
    (result_dir / "attachment1_three_method_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
