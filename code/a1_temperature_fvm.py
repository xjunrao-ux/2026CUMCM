"""A problem, question 1: radial temperature field by finite volume method.

The herb is approximated as a long axisymmetric cylinder.  The script solves
the transient radial heat-conduction equation with a time-varying convective
boundary condition supplied by Attachment 1.

Numerical method
----------------
* Node-centred finite volumes in the radial direction.
* Backward Euler time stepping.
* A reusable Thomas factorisation for the constant tridiagonal system.

Outputs
-------
* Full 1 s x 0.1 cm temperature field as CSV.
* The values requested by Table 1 as CSV and a rendered table.
* Radial-profile and temperature-field figures in SVG/PDF/PNG.
* A JSON report containing conservation, bounds, monotonicity and grid-
  refinement checks.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import dataclass
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np
from openpyxl import load_workbook


# Required figure settings: keep vector text editable.
plt.rcParams['font.family'] = 'sans-serif'
plt.rcParams['font.sans-serif'] = ['Arial', 'DejaVu Sans', 'Liberation Sans']
plt.rcParams['svg.fonttype'] = 'none'
plt.rcParams['pdf.fonttype'] = 42
mpl.rcParams.update({"svg.fonttype": "none", "pdf.fonttype": 42})


RADIUS_M = 0.02
INITIAL_TEMPERATURE_C = 28.0
DENSITY_KG_M3 = 820.0
SPECIFIC_HEAT_J_KG_K = 2600.0
THERMAL_CONDUCTIVITY_W_M_K = 0.36
CONVECTION_COEFFICIENT_W_M2_K = 25.0
END_TIME_S = 1800.0

TABLE_TIMES_S = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0], dtype=float)


@dataclass
class SimulationResult:
    times_s: np.ndarray
    radii_m: np.ndarray
    temperature_c: np.ndarray
    capacities_j_k: np.ndarray
    boundary_heat_input_j_m: float


def configure_fonts() -> None:
    """Prefer a Windows CJK font when present, with safe fallbacks."""
    available = {font.name for font in fm.fontManager.ttflist}
    candidates = ["Microsoft YaHei", "SimHei", "Arial", "DejaVu Sans"]
    selected = [name for name in candidates if name in available]
    plt.rcParams["font.sans-serif"] = selected or ["DejaVu Sans"]
    plt.rcParams.update(
        {
            "font.size": 8,
            "axes.labelsize": 8,
            "axes.titlesize": 9,
            "xtick.labelsize": 7,
            "ytick.labelsize": 7,
            "legend.fontsize": 7,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.8,
            "legend.frameon": False,
        }
    )


def find_default_data_path(repo_root: Path) -> Path:
    path = (
        repo_root
        / "比赛题目"
        / "CUMCM2026Problems"
        / "A题"
        / "附件"
        / "附件1.xlsx"
    )
    if not path.exists():
        raise FileNotFoundError(f"Attachment 1 was not found: {path}")
    return path


def load_oven_temperature(path: Path) -> tuple[np.ndarray, np.ndarray]:
    """Read time and oven temperature from Attachment 1."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    workbook.close()

    times = np.asarray([row[0] for row in rows], dtype=float)
    temperatures = np.asarray([row[1] for row in rows], dtype=float)

    if times.size < 2 or temperatures.size != times.size:
        raise ValueError("Attachment 1 does not contain a valid temperature series.")
    if not np.all(np.isfinite(times)) or not np.all(np.isfinite(temperatures)):
        raise ValueError("Attachment 1 contains missing or non-finite values.")
    if not np.all(np.diff(times) > 0):
        raise ValueError("Attachment 1 time values must be strictly increasing.")
    if times[0] > 0 or times[-1] < END_TIME_S:
        raise ValueError("Attachment 1 does not cover the required 0-1800 s interval.")
    return times, temperatures


def factor_tridiagonal(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Pre-factor a tridiagonal matrix for repeated Thomas solves."""
    n = diagonal.size
    denominators = np.empty(n, dtype=float)
    upper_prime = np.empty(n - 1, dtype=float)
    denominators[0] = diagonal[0]
    if abs(denominators[0]) < 1e-15:
        raise np.linalg.LinAlgError("Zero pivot in tridiagonal factorisation.")
    upper_prime[0] = upper[0] / denominators[0]
    for i in range(1, n - 1):
        denominators[i] = diagonal[i] - lower[i - 1] * upper_prime[i - 1]
        if abs(denominators[i]) < 1e-15:
            raise np.linalg.LinAlgError("Zero pivot in tridiagonal factorisation.")
        upper_prime[i] = upper[i] / denominators[i]
    denominators[-1] = diagonal[-1] - lower[-1] * upper_prime[-1]
    if abs(denominators[-1]) < 1e-15:
        raise np.linalg.LinAlgError("Zero pivot in tridiagonal factorisation.")
    return lower.copy(), denominators, upper_prime


def solve_factored_tridiagonal(
    factor: tuple[np.ndarray, np.ndarray, np.ndarray], rhs: np.ndarray
) -> np.ndarray:
    """Solve a tridiagonal system using a precomputed Thomas factor."""
    lower, denominators, upper_prime = factor
    n = rhs.size
    rhs_prime = np.empty(n, dtype=float)
    rhs_prime[0] = rhs[0] / denominators[0]
    for i in range(1, n):
        rhs_prime[i] = (
            rhs[i] - lower[i - 1] * rhs_prime[i - 1]
        ) / denominators[i]

    solution = np.empty(n, dtype=float)
    solution[-1] = rhs_prime[-1]
    for i in range(n - 2, -1, -1):
        solution[i] = rhs_prime[i] - upper_prime[i] * solution[i + 1]
    return solution


def build_fvm_system(
    radial_step_m: float, time_step_s: float
) -> tuple[np.ndarray, tuple[np.ndarray, np.ndarray, np.ndarray], float]:
    """Build the backward-Euler system for node-centred cylindrical FVM."""
    intervals = int(round(RADIUS_M / radial_step_m))
    if not math.isclose(intervals * radial_step_m, RADIUS_M, abs_tol=1e-12):
        raise ValueError("The radial step must divide the 0.02 m radius exactly.")

    radii = np.linspace(0.0, RADIUS_M, intervals + 1)
    n = radii.size
    west_faces = np.maximum(0.0, radii - radial_step_m / 2.0)
    east_faces = np.minimum(RADIUS_M, radii + radial_step_m / 2.0)

    # Per unit cylinder length; the common length factor cancels in dT/dt.
    volumes = math.pi * (east_faces**2 - west_faces**2)
    capacities = DENSITY_KG_M3 * SPECIFIC_HEAT_J_KG_K * volumes

    lower = np.zeros(n - 1, dtype=float)
    diagonal = np.ones(n, dtype=float)
    upper = np.zeros(n - 1, dtype=float)

    for i in range(n):
        if i > 0:
            conductance_w = (
                THERMAL_CONDUCTIVITY_W_M_K
                * 2.0
                * math.pi
                * west_faces[i]
                / radial_step_m
            )
            coefficient_w = conductance_w / capacities[i]
            diagonal[i] += time_step_s * coefficient_w
            lower[i - 1] = -time_step_s * coefficient_w

        if i < n - 1:
            conductance_e = (
                THERMAL_CONDUCTIVITY_W_M_K
                * 2.0
                * math.pi
                * east_faces[i]
                / radial_step_m
            )
            coefficient_e = conductance_e / capacities[i]
            diagonal[i] += time_step_s * coefficient_e
            upper[i] = -time_step_s * coefficient_e

    surface_area_per_length = 2.0 * math.pi * RADIUS_M
    convection_coefficient = (
        CONVECTION_COEFFICIENT_W_M2_K
        * surface_area_per_length
        / capacities[-1]
    )
    diagonal[-1] += time_step_s * convection_coefficient

    factor = factor_tridiagonal(lower, diagonal, upper)
    return radii, factor, convection_coefficient


def simulate_temperature(
    source_times_s: np.ndarray,
    source_temperature_c: np.ndarray,
    radial_step_cm: float,
    time_step_s: float,
) -> SimulationResult:
    """Run the implicit finite-volume temperature simulation."""
    radial_step_m = radial_step_cm / 100.0
    steps = int(round(END_TIME_S / time_step_s))
    if not math.isclose(steps * time_step_s, END_TIME_S, abs_tol=1e-12):
        raise ValueError("The time step must divide 1800 s exactly.")
    record_stride = int(round(1.0 / time_step_s))
    if not math.isclose(record_stride * time_step_s, 1.0, abs_tol=1e-12):
        raise ValueError("The time step must divide the 1 s output interval exactly.")

    radii, factor, convection_coefficient = build_fvm_system(
        radial_step_m, time_step_s
    )

    # Reconstruct control-volume heat capacities for energy-balance checking.
    west_faces = np.maximum(0.0, radii - radial_step_m / 2.0)
    east_faces = np.minimum(RADIUS_M, radii + radial_step_m / 2.0)
    volumes = math.pi * (east_faces**2 - west_faces**2)
    capacities = DENSITY_KG_M3 * SPECIFIC_HEAT_J_KG_K * volumes

    output_times = np.arange(0.0, END_TIME_S + 1.0, 1.0)
    temperature = np.full(radii.size, INITIAL_TEMPERATURE_C, dtype=float)
    records = np.empty((output_times.size, radii.size), dtype=float)
    records[0] = temperature

    boundary_heat_input = 0.0
    surface_area_per_length = 2.0 * math.pi * RADIUS_M
    record_index = 1

    for step in range(1, steps + 1):
        time_now = step * time_step_s
        # np.interp is valid because load_oven_temperature asserts strict order.
        oven_temperature = float(
            np.interp(time_now, source_times_s, source_temperature_c)
        )
        rhs = temperature.copy()
        rhs[-1] += time_step_s * convection_coefficient * oven_temperature
        temperature = solve_factored_tridiagonal(factor, rhs)

        boundary_heat_input += (
            time_step_s
            * CONVECTION_COEFFICIENT_W_M2_K
            * surface_area_per_length
            * (oven_temperature - temperature[-1])
        )

        if step % record_stride == 0:
            records[record_index] = temperature
            record_index += 1

    if record_index != output_times.size:
        raise RuntimeError("The number of recorded output times is inconsistent.")

    return SimulationResult(
        times_s=output_times,
        radii_m=radii,
        temperature_c=records,
        capacities_j_k=capacities,
        boundary_heat_input_j_m=boundary_heat_input,
    )


def sample_field(
    result: SimulationResult, query_times_s: np.ndarray, query_radii_cm: np.ndarray
) -> np.ndarray:
    """Sample a result on requested times and physical radii."""
    sampled = np.empty((query_times_s.size, query_radii_cm.size), dtype=float)
    radii_cm = result.radii_m * 100.0
    for row, time_s in enumerate(query_times_s):
        time_index = int(round(float(time_s)))
        sampled[row] = np.interp(
            query_radii_cm, radii_cm, result.temperature_c[time_index]
        )
    return sampled


def write_full_csv(
    path: Path, result: SimulationResult, output_radii_cm: np.ndarray
) -> None:
    source_radii_cm = result.radii_m * 100.0
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", *[f"{r:.1f}_cm" for r in output_radii_cm]])
        for time_s, row in zip(result.times_s, result.temperature_c):
            output_row = np.interp(output_radii_cm, source_radii_cm, row)
            writer.writerow([int(time_s), *[f"{value:.6f}" for value in output_row]])


def write_table_csv(path: Path, table_values: np.ndarray) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", *[f"{r:g}_cm" for r in TABLE_RADII_CM]])
        for time_s, row in zip(TABLE_TIMES_S, table_values):
            writer.writerow([int(time_s), *[f"{value:.4f}" for value in row]])


def save_figure(fig: plt.Figure, stem: Path) -> None:
    fig.savefig(stem.with_suffix(".svg"), bbox_inches="tight")
    fig.savefig(stem.with_suffix(".pdf"), bbox_inches="tight")
    fig.savefig(stem.with_suffix(".png"), dpi=600, bbox_inches="tight")
    fig.savefig(stem.with_suffix(".tiff"), dpi=600, bbox_inches="tight")
    plt.close(fig)


def plot_profiles(figure_dir: Path, result: SimulationResult) -> None:
    fig, ax = plt.subplots(figsize=(7.09, 4.13))
    colors = plt.cm.viridis(np.linspace(0.12, 0.92, TABLE_TIMES_S.size))
    radii_cm = result.radii_m * 100.0
    for time_s, color in zip(TABLE_TIMES_S, colors):
        row = result.temperature_c[int(time_s)]
        ax.plot(
            radii_cm,
            row,
            color=color,
            linewidth=1.7,
            marker="o",
            markersize=2.5,
            label=f"{int(time_s)} s",
        )
    ax.set_xlabel("到药材中心的距离 (cm)")
    ax.set_ylabel("温度 (°C)")
    ax.set_title("30分钟内药材径向温度分布")
    ax.set_xlim(0.0, 2.0)
    ax.grid(axis="both", color="#D9D9D9", linewidth=0.5, alpha=0.8)
    ax.legend(ncol=2, loc="upper left")
    save_figure(fig, figure_dir / "temperature_profiles")


def plot_heatmap(figure_dir: Path, result: SimulationResult) -> None:
    fig, ax = plt.subplots(figsize=(7.09, 3.94))
    image = ax.imshow(
        result.temperature_c.T,
        origin="lower",
        aspect="auto",
        interpolation="nearest",
        extent=[0.0, END_TIME_S / 60.0, 0.0, RADIUS_M * 100.0],
        cmap="YlOrRd",
    )
    colorbar = fig.colorbar(image, ax=ax, pad=0.02)
    colorbar.set_label("温度 (°C)")
    ax.set_xlabel("时间 (min)")
    ax.set_ylabel("到药材中心的距离 (cm)")
    ax.set_title("药材内部温度的时空变化")
    save_figure(fig, figure_dir / "temperature_heatmap")


def plot_table(figure_dir: Path, table_values: np.ndarray) -> None:
    column_labels = [f"{radius:g}" for radius in TABLE_RADII_CM]
    row_labels = [f"{int(time_s)}" for time_s in TABLE_TIMES_S]
    cell_text = [[f"{value:.4f}" for value in row] for row in table_values]

    fig, ax = plt.subplots(figsize=(6.69, 3.07))
    ax.axis("off")
    ax.set_title("表1  30分钟内药材的温度", fontsize=10, fontweight="bold", pad=9)
    table = ax.table(
        cellText=cell_text,
        rowLabels=row_labels,
        colLabels=column_labels,
        cellLoc="center",
        rowLoc="center",
        loc="center",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8)
    table.scale(1.0, 1.35)

    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("#7A7A7A")
        cell.set_linewidth(0.6)
        if row == 0:
            cell.set_facecolor("#DDE8F2")
            cell.set_text_props(weight="bold")
        elif col == -1:
            cell.set_facecolor("#F1F3F5")
            cell.set_text_props(weight="bold")

    ax.text(
        0.5,
        0.99,
        "到药材中心的距离 / cm",
        transform=ax.transAxes,
        ha="center",
        va="top",
        fontsize=8,
    )
    ax.text(
        -0.015,
        0.83,
        "时间 / s",
        transform=ax.transAxes,
        ha="right",
        va="center",
        fontsize=8,
    )
    save_figure(fig, figure_dir / "table1_temperature")


def build_validation_report(
    baseline: SimulationResult,
    medium: SimulationResult,
    production: SimulationResult,
    reference: SimulationResult,
    source_times_s: np.ndarray,
    source_temperature_c: np.ndarray,
) -> dict:
    baseline_table = sample_field(baseline, TABLE_TIMES_S, TABLE_RADII_CM)
    medium_table = sample_field(medium, TABLE_TIMES_S, TABLE_RADII_CM)
    production_table = sample_field(production, TABLE_TIMES_S, TABLE_RADII_CM)
    reference_table = sample_field(reference, TABLE_TIMES_S, TABLE_RADII_CM)

    baseline_error = np.abs(baseline_table - reference_table)
    medium_error = np.abs(medium_table - reference_table)
    production_error = np.abs(production_table - reference_table)

    energy_change = float(
        np.sum(
            production.capacities_j_k
            * (production.temperature_c[-1] - INITIAL_TEMPERATURE_C)
        )
    )
    energy_residual = energy_change - production.boundary_heat_input_j_m
    energy_scale = max(abs(production.boundary_heat_input_j_m), 1.0)

    oven_at_seconds = np.interp(
        production.times_s, source_times_s, source_temperature_c
    )
    radial_differences = np.diff(production.temperature_c, axis=1)
    temporal_differences = np.diff(production.temperature_c, axis=0)

    alpha = THERMAL_CONDUCTIVITY_W_M_K / (
        DENSITY_KG_M3 * SPECIFIC_HEAT_J_KG_K
    )
    biot = CONVECTION_COEFFICIENT_W_M2_K * RADIUS_M / THERMAL_CONDUCTIVITY_W_M_K
    fourier = alpha * END_TIME_S / (RADIUS_M**2)

    return {
        "model": {
            "geometry": "one-dimensional axisymmetric cylinder",
            "radial_method": "node-centred finite volume",
            "time_method": "backward Euler",
            "output_radial_interval_cm": 0.1,
            "output_time_interval_s": 1.0,
            "production_radial_step_cm": 0.025,
            "production_time_step_s": 0.25,
            "reference_radial_step_cm": 0.0125,
            "reference_time_step_s": 0.125,
        },
        "dimensionless_numbers": {
            "thermal_diffusivity_m2_s": alpha,
            "Biot_number": biot,
            "Fourier_number_at_1800_s": fourier,
        },
        "grid_refinement": {
            "baseline_vs_reference_max_abs_C": float(np.max(baseline_error)),
            "baseline_vs_reference_mean_abs_C": float(np.mean(baseline_error)),
            "medium_vs_reference_max_abs_C": float(np.max(medium_error)),
            "medium_vs_reference_mean_abs_C": float(np.mean(medium_error)),
            "production_vs_reference_max_abs_C": float(np.max(production_error)),
            "production_vs_reference_mean_abs_C": float(np.mean(production_error)),
        },
        "energy_balance_per_unit_length": {
            "stored_energy_change_J_m": energy_change,
            "integrated_convective_input_J_m": production.boundary_heat_input_j_m,
            "residual_J_m": energy_residual,
            "relative_residual": energy_residual / energy_scale,
        },
        "physical_checks": {
            "minimum_predicted_temperature_C": float(np.min(production.temperature_c)),
            "maximum_predicted_temperature_C": float(np.max(production.temperature_c)),
            "maximum_oven_temperature_0_1800_s_C": float(np.max(oven_at_seconds)),
            "radial_order_violation_count": int(np.sum(radial_differences < -1e-10)),
            "time_monotonicity_violation_count": int(
                np.sum(temporal_differences < -1e-10)
            ),
            "temperature_above_oven_violation_count": int(
                np.sum(production.temperature_c[:, -1] - oven_at_seconds > 1e-10)
            ),
        },
        "selected_outputs": {
            "center_temperature_1800_s_C": float(production.temperature_c[-1, 0]),
            "surface_temperature_1800_s_C": float(production.temperature_c[-1, -1]),
            "center_surface_gap_1800_s_C": float(
                production.temperature_c[-1, -1] - production.temperature_c[-1, 0]
            ),
            "oven_temperature_1800_s_C": float(oven_at_seconds[-1]),
        },
    }


def print_table(table_values: np.ndarray) -> None:
    headers = ["time/s", *[f"{r:g} cm" for r in TABLE_RADII_CM]]
    widths = [9, 11, 11, 11, 11, 11]
    line = " ".join(f"{header:>{width}}" for header, width in zip(headers, widths))
    print(line)
    print("-" * len(line))
    for time_s, row in zip(TABLE_TIMES_S, table_values):
        values = [f"{int(time_s)}", *[f"{value:.4f}" for value in row]]
        print(" ".join(f"{value:>{width}}" for value, width in zip(values, widths)))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data",
        type=Path,
        default=None,
        help="Path to Attachment 1. Defaults to the official file in the repository.",
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=None,
        help="Directory for CSV and validation outputs.",
    )
    parser.add_argument(
        "--figure-dir",
        type=Path,
        default=None,
        help="Directory for SVG, PDF and PNG figures.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    repo_root = Path(__file__).resolve().parents[1]
    data_path = args.data or find_default_data_path(repo_root)
    results_dir = args.results_dir or repo_root / "results" / "A_problem1_temperature"
    figure_dir = args.figure_dir or repo_root / "picture" / "A_problem1_temperature"
    results_dir.mkdir(parents=True, exist_ok=True)
    figure_dir.mkdir(parents=True, exist_ok=True)

    source_times_s, source_temperature_c = load_oven_temperature(data_path)

    baseline = simulate_temperature(
        source_times_s, source_temperature_c, radial_step_cm=0.1, time_step_s=1.0
    )
    medium = simulate_temperature(
        source_times_s, source_temperature_c, radial_step_cm=0.05, time_step_s=0.5
    )
    production = simulate_temperature(
        source_times_s, source_temperature_c, radial_step_cm=0.025, time_step_s=0.25
    )
    reference = simulate_temperature(
        source_times_s,
        source_temperature_c,
        radial_step_cm=0.0125,
        time_step_s=0.125,
    )

    table_values = sample_field(production, TABLE_TIMES_S, TABLE_RADII_CM)
    report = build_validation_report(
        baseline,
        medium,
        production,
        reference,
        source_times_s,
        source_temperature_c,
    )

    output_radii_cm = np.arange(0.0, 2.0 + 0.1, 0.1)
    write_full_csv(
        results_dir / "temperature_full_1s_0p1cm.csv",
        production,
        output_radii_cm,
    )
    write_table_csv(results_dir / "table1_temperature.csv", table_values)
    with (results_dir / "validation_summary.json").open(
        "w", encoding="utf-8"
    ) as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)

    configure_fonts()
    plot_profiles(figure_dir, production)
    plot_heatmap(figure_dir, production)
    plot_table(figure_dir, table_values)

    print("Table 1: radial temperature in the first 30 minutes (deg C)")
    print_table(table_values)
    print("\nValidation summary")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"\nResults: {results_dir}")
    print(f"Figures: {figure_dir}")


if __name__ == "__main__":
    main()
