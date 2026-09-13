"""A题附件1、附件2预处理：分段线性插值与单位换算。

默认输入：
    2026A/比赛题目/CUMCM2026Problems/A题/附件/附件1.xlsx
    2026A/比赛题目/CUMCM2026Problems/A题/附件/附件2.xlsx

默认输出：
    2026A/results/A_preprocessing/附件1_1s线性插值.xlsx
    2026A/results/A_preprocessing/附件2_60s线性插值.xlsx

依赖：openpyxl
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


ATTACHMENT1_SOURCE_STEP_S = 60
ATTACHMENT2_SOURCE_STEP_S = 1800
ATTACHMENT1_OUTPUT_STEP_S = 1
ATTACHMENT2_OUTPUT_STEP_S = 60


def project_root() -> Path:
    """返回 2026A 项目目录。"""
    return Path(__file__).resolve().parents[1]


def read_numeric_rows(path: Path, column_count: int) -> list[tuple[float, ...]]:
    """读取第一个工作表，跳过表头，并验证指定列均为数值。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows: list[tuple[float, ...]] = []

    for excel_row, values in enumerate(
        worksheet.iter_rows(min_row=2, max_col=column_count, values_only=True),
        start=2,
    ):
        if all(value is None for value in values):
            continue
        if any(value is None for value in values):
            workbook.close()
            raise ValueError(f"{path.name} 第 {excel_row} 行存在缺失值：{values}")
        try:
            rows.append(tuple(float(value) for value in values))
        except (TypeError, ValueError) as error:
            workbook.close()
            raise ValueError(
                f"{path.name} 第 {excel_row} 行包含非数值：{values}"
            ) from error

    workbook.close()
    if not rows:
        raise ValueError(f"{path.name} 没有有效数据")
    return rows


def validate_time_series(
    rows: Sequence[Sequence[float]], expected_source_step_s: int, file_name: str
) -> None:
    """检查时间严格递增且原始时间间隔符合题目附件。"""
    for index in range(1, len(rows)):
        delta = rows[index][0] - rows[index - 1][0]
        if delta <= 0:
            raise ValueError(f"{file_name} 时间列在第 {index + 2} 行未严格递增")
        if abs(delta - expected_source_step_s) > 1e-9:
            raise ValueError(
                f"{file_name} 第 {index + 2} 行时间间隔为 {delta} s，"
                f"预期为 {expected_source_step_s} s"
            )


def piecewise_linear_interpolation(
    rows: Sequence[Sequence[float]],
    output_step_s: int,
    value_columns: Sequence[int],
) -> list[tuple[float, ...]]:
    """仅在相邻原始观测点之间进行分段线性插值，不做外推。"""
    start_time = rows[0][0]
    end_time = rows[-1][0]
    if not start_time.is_integer() or not end_time.is_integer():
        raise ValueError("时间端点必须是整数秒")
    if (int(end_time) - int(start_time)) % output_step_s != 0:
        raise ValueError("输出步长不能整除时间范围")

    output: list[tuple[float, ...]] = []
    left_index = 0

    for time_s in range(int(start_time), int(end_time) + 1, output_step_s):
        while left_index < len(rows) - 2 and time_s > rows[left_index + 1][0]:
            left_index += 1

        left = rows[left_index]
        right = rows[min(left_index + 1, len(rows) - 1)]
        interval = right[0] - left[0]
        alpha = 0.0 if interval == 0 else (time_s - left[0]) / interval
        values = tuple(
            left[column] + alpha * (right[column] - left[column])
            for column in value_columns
        )
        output.append((float(time_s), *values))

    return output


def clean_number(value: float, digits: int = 12) -> float | int:
    """抑制浮点尾差；整数时间按整数写入 Excel。"""
    rounded = round(float(value), digits)
    if rounded.is_integer():
        return int(rounded)
    return rounded


def write_plain_workbook(
    output_path: Path,
    headers: Sequence[str],
    rows: Iterable[Sequence[float]],
    column_widths: Sequence[float],
) -> None:
    """写出无颜色、无边框和无表格样式的单工作表 Excel 文件。"""
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("插值数据")
    for column_index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append([clean_number(value) for value in row])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_saved_workbook(
    path: Path,
    expected_data_rows: int,
    expected_columns: int,
) -> tuple[tuple[object, ...], tuple[object, ...]]:
    """重新读取输出文件，检查尺寸以及首末数据行。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header = tuple(next(iterator))
    if len(header) != expected_columns:
        workbook.close()
        raise RuntimeError(
            f"{path.name} 列数错误：{len(header)}，预期 {expected_columns}"
        )

    first_row: tuple[object, ...] | None = None
    last_row: tuple[object, ...] | None = None
    data_row_count = 0
    for values in iterator:
        current_row = tuple(values[:expected_columns])
        if first_row is None:
            first_row = current_row
        last_row = current_row
        data_row_count += 1

    if data_row_count != expected_data_rows:
        workbook.close()
        raise RuntimeError(
            f"{path.name} 行数错误：{data_row_count}，预期 {expected_data_rows}"
        )
    if first_row is None or last_row is None:
        workbook.close()
        raise RuntimeError(f"{path.name} 没有数据行")
    workbook.close()
    return first_row, last_row


def parse_args() -> argparse.Namespace:
    root = project_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--attachment-dir",
        type=Path,
        default=root / "比赛题目" / "CUMCM2026Problems" / "A题" / "附件",
        help="包含附件1.xlsx和附件2.xlsx的目录",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=root / "results" / "A_preprocessing",
        help="两个结果工作簿的输出目录",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    attachment_dir = args.attachment_dir.resolve()
    output_dir = args.output_dir.resolve()
    attachment1_path = attachment_dir / "附件1.xlsx"
    attachment2_path = attachment_dir / "附件2.xlsx"

    source1 = read_numeric_rows(attachment1_path, column_count=3)
    source2 = read_numeric_rows(attachment2_path, column_count=2)
    validate_time_series(source1, ATTACHMENT1_SOURCE_STEP_S, attachment1_path.name)
    validate_time_series(source2, ATTACHMENT2_SOURCE_STEP_S, attachment2_path.name)

    # 附件1列：时间(s)、温度(°C)、水分浓度(kg/kg)。
    interpolated1 = piecewise_linear_interpolation(
        source1,
        output_step_s=ATTACHMENT1_OUTPUT_STEP_S,
        value_columns=(1, 2),
    )
    attachment1_rows = [
        (time_s, temperature_c, temperature_c + 273.15, moisture)
        for time_s, temperature_c, moisture in interpolated1
    ]

    # 附件2列：时间(s)、半径(cm)。
    interpolated2 = piecewise_linear_interpolation(
        source2,
        output_step_s=ATTACHMENT2_OUTPUT_STEP_S,
        value_columns=(1,),
    )
    attachment2_rows = [
        (time_s, radius_cm, radius_cm / 100.0)
        for time_s, radius_cm in interpolated2
    ]

    output1 = output_dir / "附件1_1s线性插值.xlsx"
    output2 = output_dir / "附件2_60s线性插值.xlsx"
    write_plain_workbook(
        output1,
        ("时间（s）", "温度（°C）", "温度（K）", "水分浓度（kg/kg）"),
        attachment1_rows,
        (14, 16, 16, 24),
    )
    write_plain_workbook(
        output2,
        ("时间（s）", "半径（cm）", "半径（m）"),
        attachment2_rows,
        (14, 16, 16),
    )

    first1, last1 = verify_saved_workbook(output1, len(attachment1_rows), 4)
    first2, last2 = verify_saved_workbook(output2, len(attachment2_rows), 3)
    print(f"已生成：{output1}")
    print(f"附件1数据行数：{len(attachment1_rows)}；首行：{first1}；末行：{last1}")
    print(f"已生成：{output2}")
    print(f"附件2数据行数：{len(attachment2_rows)}；首行：{first2}；末行：{last2}")


if __name__ == "__main__":
    main()
