"""Compare one-second piecewise-linear and PCHIP fits for Attachment 1."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.font_manager as fm
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


END_TIME_S = 1800.0


def configure_fonts() -> None:
    available = {font.name for font in fm.fontManager.ttflist}
    candidates = ["Microsoft YaHei", "SimHei", "Arial", "DejaVu Sans"]
    selected = [name for name in candidates if name in available]
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": selected or ["DejaVu Sans"],
            "axes.unicode_minus": False,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "font.size": 8.0,
            "axes.labelsize": 8.0,
            "axes.titlesize": 9.0,
            "xtick.labelsize": 7.0,
            "ytick.labelsize": 7.0,
            "legend.fontsize": 7.0,
            "axes.linewidth": 0.7,
            "legend.frameon": False,
        }
    )


def load_attachment(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = [
        row[:3]
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None and float(row[0]) <= END_TIME_S
    ]
    workbook.close()
    data = np.asarray(rows, dtype=float)
    times_s, temperature_c, moisture_kg_kg = data.T
    if times_s[0] != 0.0 or times_s[-1] != END_TIME_S:
        raise ValueError("Attachment 1 does not cover the required 0-1800 s interval.")
    if not np.all(np.diff(times_s) > 0.0):
        raise ValueError("Attachment 1 time values must be strictly increasing.")
    return times_s, temperature_c, moisture_kg_kg


def fit_one_second(
    times_s: np.ndarray,
    values: np.ndarray,
    query_times_s: np.ndarray,
) -> tuple[np.ndarray, np.ndarray]:
    linear = np.interp(query_times_s, times_s, values)
    pchip = PchipInterpolator(times_s, values)(query_times_s)
    return linear, pchip


def calculate_metrics(
    variable: str,
    unit: str,
    times_s: np.ndarray,
    observed: np.ndarray,
    query_times_s: np.ndarray,
    linear: np.ndarray,
    pchip: np.ndarray,
) -> dict[str, float | int | str]:
    difference = pchip - linear
    absolute_difference = np.abs(difference)
    maximum_index = int(np.argmax(absolute_difference))
    knot_indices = np.searchsorted(query_times_s, times_s)

    segment = np.minimum(
        np.searchsorted(times_s, query_times_s, side="right") - 1,
        len(times_s) - 2,
    )
    segment = np.maximum(segment, 0)
    lower = np.minimum(observed[segment], observed[segment + 1])
    upper = np.maximum(observed[segment], observed[segment + 1])
    overshoot_count = int(
        np.count_nonzero((pchip < lower - 1e-12) | (pchip > upper + 1e-12))
    )

    return {
        "variable": variable,
        "unit": unit,
        "observation_count": int(times_s.size),
        "observation_interval_s": float(np.median(np.diff(times_s))),
        "maximum_absolute_difference": float(absolute_difference[maximum_index]),
        "maximum_difference_time_s": float(query_times_s[maximum_index]),
        "signed_difference_at_maximum": float(difference[maximum_index]),
        "mean_absolute_difference": float(np.mean(absolute_difference)),
        "root_mean_square_difference": float(np.sqrt(np.mean(difference**2))),
        "maximum_relative_to_observed_range_percent": float(
            100.0 * absolute_difference[maximum_index] / np.ptp(observed)
        ),
        "maximum_knot_difference": float(
            np.max(np.abs(difference[knot_indices]))
        ),
        "pchip_local_overshoot_count": overshoot_count,
    }


def write_full_csv(
    path: Path,
    source_times_s: np.ndarray,
    source_temperature_c: np.ndarray,
    source_moisture_kg_kg: np.ndarray,
    query_times_s: np.ndarray,
    temperature_linear: np.ndarray,
    temperature_pchip: np.ndarray,
    moisture_linear: np.ndarray,
    moisture_pchip: np.ndarray,
) -> None:
    source_lookup = {
        int(time_s): (temperature, moisture)
        for time_s, temperature, moisture in zip(
            source_times_s, source_temperature_c, source_moisture_kg_kg
        )
    }
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "time_s",
                "is_observation",
                "observed_temperature_C",
                "temperature_linear_C",
                "temperature_pchip_C",
                "temperature_pchip_minus_linear_C",
                "temperature_absolute_difference_C",
                "observed_moisture_kg_kg",
                "moisture_linear_kg_kg",
                "moisture_pchip_kg_kg",
                "moisture_pchip_minus_linear_kg_kg",
                "moisture_absolute_difference_kg_kg",
            ]
        )
        for index, time_s in enumerate(query_times_s.astype(int)):
            source = source_lookup.get(time_s)
            temperature_difference = temperature_pchip[index] - temperature_linear[index]
            moisture_difference = moisture_pchip[index] - moisture_linear[index]
            writer.writerow(
                [
                    time_s,
                    int(source is not None),
                    "" if source is None else f"{source[0]:.6f}",
                    f"{temperature_linear[index]:.10f}",
                    f"{temperature_pchip[index]:.10f}",
                    f"{temperature_difference:.10f}",
                    f"{abs(temperature_difference):.10f}",
                    "" if source is None else f"{source[1]:.8f}",
                    f"{moisture_linear[index]:.12f}",
                    f"{moisture_pchip[index]:.12f}",
                    f"{moisture_difference:.12f}",
                    f"{abs(moisture_difference):.12f}",
                ]
            )


def write_summary_csv(path: Path, metrics: list[dict]) -> None:
    headers = list(metrics[0])
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(metrics)


def style_axis(axis: plt.Axes) -> None:
    axis.grid(True, color="#D9D9D9", linewidth=0.45, alpha=0.75)
    axis.set_axisbelow(True)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)


def render_comparison_png(
    path: Path,
    source_times_s: np.ndarray,
    source_temperature_c: np.ndarray,
    source_moisture_kg_kg: np.ndarray,
    query_times_s: np.ndarray,
    temperature_linear: np.ndarray,
    temperature_pchip: np.ndarray,
    moisture_linear: np.ndarray,
    moisture_pchip: np.ndarray,
) -> None:
    configure_fonts()
    time_minutes = query_times_s / 60.0
    source_minutes = source_times_s / 60.0
    figure, axes = plt.subplots(
        2, 2, figsize=(7.09, 5.35), constrained_layout=True,
        gridspec_kw={"height_ratios": [1.45, 1.0]},
    )
    colors = {"linear": "#0072B2", "pchip": "#D55E00", "difference": "#7B3294"}

    axes[0, 0].plot(
        time_minutes, temperature_linear, color=colors["linear"], linewidth=1.5,
        label="分段线性",
    )
    axes[0, 0].plot(
        time_minutes, temperature_pchip, color=colors["pchip"], linewidth=1.25,
        linestyle="--", label="PCHIP",
    )
    axes[0, 0].scatter(
        source_minutes, source_temperature_c, s=10, color="#222222", zorder=3,
        label="附件1观测点",
    )
    axes[0, 0].set(
        title="(a) 烘房温度拟合",
        ylabel="温度 / °C",
        xlim=(0.0, 30.0),
    )

    axes[0, 1].plot(
        time_minutes, moisture_linear, color=colors["linear"], linewidth=1.5,
        label="分段线性",
    )
    axes[0, 1].plot(
        time_minutes, moisture_pchip, color=colors["pchip"], linewidth=1.25,
        linestyle="--", label="PCHIP",
    )
    axes[0, 1].scatter(
        source_minutes, source_moisture_kg_kg, s=10, color="#222222", zorder=3,
        label="附件1观测点",
    )
    axes[0, 1].set(
        title="(b) 烘房水分浓度拟合",
        ylabel="水分浓度 / (kg/kg)",
        xlim=(0.0, 30.0),
    )

    temperature_difference = temperature_pchip - temperature_linear
    moisture_difference = moisture_pchip - moisture_linear
    axes[1, 0].plot(
        time_minutes, temperature_difference, color=colors["difference"], linewidth=1.15
    )
    axes[1, 0].fill_between(
        time_minutes, 0.0, temperature_difference,
        color=colors["difference"], alpha=0.16, linewidth=0,
    )
    axes[1, 0].set(
        title="(c) 温度拟合差值（PCHIP−线性）",
        xlabel="时间 / min",
        ylabel="差值 / °C",
        xlim=(0.0, 30.0),
    )

    axes[1, 1].plot(
        time_minutes, moisture_difference, color=colors["difference"], linewidth=1.15
    )
    axes[1, 1].fill_between(
        time_minutes, 0.0, moisture_difference,
        color=colors["difference"], alpha=0.16, linewidth=0,
    )
    axes[1, 1].set(
        title="(d) 水分拟合差值（PCHIP−线性）",
        xlabel="时间 / min",
        ylabel="差值 / (kg/kg)",
        xlim=(0.0, 30.0),
    )

    for axis in axes.ravel():
        style_axis(axis)
        axis.set_xticks(np.arange(0.0, 31.0, 5.0))
    axes[1, 0].axhline(0.0, color="#555555", linewidth=0.65)
    axes[1, 1].axhline(0.0, color="#555555", linewidth=0.65)
    axes[0, 0].legend(loc="lower right", ncol=1)
    figure.savefig(path, dpi=300, facecolor="white")
    plt.close(figure)


def render_summary_table_png(path: Path, metrics: list[dict]) -> None:
    configure_fonts()
    figure, axis = plt.subplots(figsize=(7.09, 2.25))
    axis.axis("off")
    axis.set_title("附件1分段线性与PCHIP拟合差异汇总", fontsize=11, fontweight="bold", pad=14)
    columns = ["变量", "最大绝对差", "平均绝对差", "均方根差", "最大差时刻/s", "局部过冲数"]
    rows = []
    for item in metrics:
        if item["variable"] == "temperature":
            formats = ["温度", ".6f", ".6f", ".6f"]
        else:
            formats = ["水分浓度", ".8f", ".8f", ".8f"]
        rows.append(
            [
                formats[0],
                f"{item['maximum_absolute_difference']:{formats[1]}}",
                f"{item['mean_absolute_difference']:{formats[2]}}",
                f"{item['root_mean_square_difference']:{formats[3]}}",
                f"{item['maximum_difference_time_s']:.0f}",
                f"{item['pchip_local_overshoot_count']}",
            ]
        )
    table = axis.table(
        cellText=rows,
        colLabels=columns,
        cellLoc="center",
        colLoc="center",
        bbox=[0.02, 0.23, 0.96, 0.58],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(7.5)
    for (row, _column), cell in table.get_celld().items():
        cell.set_edgecolor("#8A8A8A")
        cell.set_linewidth(0.7)
        if row == 0:
            cell.set_facecolor("#D9E4F0")
            cell.set_text_props(fontweight="bold")
    axis.text(
        0.02, 0.11,
        "注：差值定义为PCHIP拟合值减去分段线性拟合值；温度单位为°C，水分浓度单位为kg/kg。",
        transform=axis.transAxes, fontsize=7.0, color="#555555", ha="left",
    )
    figure.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()

    result_dir = args.output_root / "results" / "A_problem1_interpolation"
    picture_dir = args.output_root / "picture" / "A_problem1_interpolation"
    result_dir.mkdir(parents=True, exist_ok=True)
    picture_dir.mkdir(parents=True, exist_ok=True)

    times_s, temperature_c, moisture_kg_kg = load_attachment(args.data)
    query_times_s = np.arange(0.0, END_TIME_S + 1.0)
    temperature_linear, temperature_pchip = fit_one_second(
        times_s, temperature_c, query_times_s
    )
    moisture_linear, moisture_pchip = fit_one_second(
        times_s, moisture_kg_kg, query_times_s
    )
    metrics = [
        calculate_metrics(
            "temperature", "degC", times_s, temperature_c, query_times_s,
            temperature_linear, temperature_pchip,
        ),
        calculate_metrics(
            "moisture", "kg/kg", times_s, moisture_kg_kg, query_times_s,
            moisture_linear, moisture_pchip,
        ),
    ]

    write_full_csv(
        result_dir / "attachment1_interpolation_comparison_1s.csv",
        times_s, temperature_c, moisture_kg_kg, query_times_s,
        temperature_linear, temperature_pchip, moisture_linear, moisture_pchip,
    )
    write_summary_csv(
        result_dir / "attachment1_interpolation_summary.csv", metrics
    )
    render_comparison_png(
        picture_dir / "attachment1_interpolation_comparison.png",
        times_s, temperature_c, moisture_kg_kg, query_times_s,
        temperature_linear, temperature_pchip, moisture_linear, moisture_pchip,
    )
    render_summary_table_png(
        picture_dir / "attachment1_interpolation_summary_table.png", metrics
    )

    print(json.dumps(metrics, ensure_ascii=False, indent=2))
    print(result_dir)
    print(picture_dir)


if __name__ == "__main__":
    main()
