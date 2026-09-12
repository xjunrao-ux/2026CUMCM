"""A problem, question 2: nonlinear coupled heat-moisture finite volumes.

The whole 0-3 h process uses the Appendix 3 constitutive laws from t=0.
Oven temperature and moisture are linearly interpolated from Attachment 1.
Temperature is stored in degrees Celsius but converted to kelvin when the
Arrhenius moisture diffusivity is evaluated.  Evaporative latent heat and
shrinkage are deliberately excluded in question 2.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


RADIUS_M = 0.02
END_TIME_S = 10_800.0
INITIAL_TEMPERATURE_C = 28.0
INITIAL_MOISTURE_KG_KG = 2.55
HEAT_TRANSFER_COEFFICIENT_W_M2_K = 25.0
MASS_TRANSFER_COEFFICIENT_M_S = 8.0e-7

TABLE_TIMES_S = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
TABLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0], dtype=float)
OUTPUT_RADII_CM = np.round(np.arange(0.0, 2.0 + 0.05, 0.1), 1)


@dataclass
class EnvironmentData:
    source_times_s: np.ndarray
    source_temperature_c: np.ndarray
    source_moisture_kg_kg: np.ndarray
    times_s: np.ndarray
    temperature_c: np.ndarray
    temperature_k: np.ndarray
    moisture_kg_kg: np.ndarray


@dataclass
class Grid:
    radii_m: np.ndarray
    west_faces_m: np.ndarray
    east_faces_m: np.ndarray
    volumes_m3_m: np.ndarray


@dataclass
class CoupledResult:
    times_s: np.ndarray
    grid: Grid
    temperature_c: np.ndarray
    moisture_kg_kg: np.ndarray
    surface_temperature_c: np.ndarray
    surface_moisture_kg_kg: np.ndarray
    integrated_heat_inflow_w_s_m: float
    integrated_moisture_outflow_m3_kg_kg_m: float
    maximum_picard_iterations: int


def find_default_data_path(repo_root: Path) -> Path:
    path = (
        repo_root
        / "比赛题目"
        / "CUMCM2026Problems"
        / "A题"
        / "附件"
        / "附件1.xlsx"
    )
    if not path.exists():
        raise FileNotFoundError(f"Cannot find Attachment 1: {path}")
    return path


def load_and_preprocess_environment(path: Path) -> EnvironmentData:
    """Read Attachment 1 and interpolate its 0-3 h records to every second."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [row[:3] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    workbook.close()

    source_times = np.asarray([float(row[0]) for row in rows], dtype=float)
    source_temperature = np.asarray([float(row[1]) for row in rows], dtype=float)
    source_moisture = np.asarray([float(row[2]) for row in rows], dtype=float)
    if not np.all(np.diff(source_times) > 0.0):
        raise ValueError("Attachment 1 time values must be strictly increasing.")
    if source_times[0] > 0.0 or source_times[-1] < END_TIME_S:
        raise ValueError("Attachment 1 does not cover the required 0-10800 s interval.")
    if not (
        np.all(np.isfinite(source_temperature))
        and np.all(np.isfinite(source_moisture))
    ):
        raise ValueError("Attachment 1 contains a non-finite environmental value.")

    times = np.arange(0.0, END_TIME_S + 1.0, 1.0)
    temperature_c = np.interp(times, source_times, source_temperature)
    moisture = np.interp(times, source_times, source_moisture)
    temperature_k = temperature_c + 273.15
    return EnvironmentData(
        source_times_s=source_times,
        source_temperature_c=source_temperature,
        source_moisture_kg_kg=source_moisture,
        times_s=times,
        temperature_c=temperature_c,
        temperature_k=temperature_k,
        moisture_kg_kg=moisture,
    )


def density_kg_m3(moisture: np.ndarray) -> np.ndarray:
    return 650.0 + 128.0 * moisture


def heat_capacity_j_kg_k(moisture: np.ndarray) -> np.ndarray:
    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def conductivity_w_m_k(moisture: np.ndarray) -> np.ndarray:
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity_m2_s(
    moisture: np.ndarray, temperature_c: np.ndarray
) -> np.ndarray:
    """Appendix 3: D=2.4e-3 exp(-0.45/C) exp(-3850/T), T in K."""
    moisture = np.asarray(moisture, dtype=float)
    temperature_k = np.asarray(temperature_c, dtype=float) + 273.15
    if np.any(moisture <= 0.0) or np.any(temperature_k <= 0.0):
        raise ValueError("Positive moisture and absolute temperature are required.")
    return (
        2.4e-3
        * np.exp(-0.45 / moisture)
        * np.exp(-3850.0 / temperature_k)
    )


def make_grid(nominal_radial_step_cm: float) -> Grid:
    radial_step_m = nominal_radial_step_cm / 100.0
    intervals = int(round(RADIUS_M / radial_step_m))
    if not math.isclose(intervals * radial_step_m, RADIUS_M, abs_tol=1e-12):
        raise ValueError("The nominal radial step must divide the radius exactly.")

    logical_faces = np.linspace(0.0, 1.0, intervals + 1)
    faces = RADIUS_M * (1.0 - (1.0 - logical_faces) ** 2)
    west_faces = faces[:-1]
    east_faces = faces[1:]
    radii = 0.5 * (west_faces + east_faces)
    volumes = math.pi * (east_faces**2 - west_faces**2)
    return Grid(radii, west_faces, east_faces, volumes)


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    rhs: np.ndarray,
) -> np.ndarray:
    """Thomas algorithm for a strictly diagonally dominant system."""
    n = diagonal.size
    d = diagonal.copy()
    b = rhs.copy()
    u = upper.copy()
    for index in range(1, n):
        multiplier = lower[index - 1] / d[index - 1]
        d[index] -= multiplier * u[index - 1]
        b[index] -= multiplier * b[index - 1]
    if np.any(d <= 0.0):
        raise ValueError("The tridiagonal system lost positive definiteness.")
    solution = np.empty(n, dtype=float)
    solution[-1] = b[-1] / d[-1]
    for index in range(n - 2, -1, -1):
        solution[index] = (b[index] - u[index] * solution[index + 1]) / d[index]
    return solution


def internal_conductances(grid: Grid, coefficient: np.ndarray) -> np.ndarray:
    """Conductance at each internal face using two half-cell resistances."""
    face_radii = grid.east_faces_m[:-1]
    resistance = (
        (face_radii - grid.radii_m[:-1]) / coefficient[:-1]
        + (grid.radii_m[1:] - face_radii) / coefficient[1:]
    )
    return 2.0 * math.pi * face_radii / resistance


def build_diffusion_system(
    grid: Grid,
    field_previous: np.ndarray,
    diffusion_coefficient: np.ndarray,
    storage_coefficient: np.ndarray,
    boundary_transfer_coefficient: float,
    boundary_value: float,
    time_step_s: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, float]:
    """Assemble one backward-Euler cylindrical finite-volume system."""
    n = grid.radii_m.size
    conductances = internal_conductances(grid, diffusion_coefficient)
    storage = storage_coefficient * grid.volumes_m3_m / time_step_s
    lower = -conductances.copy()
    upper = -conductances.copy()
    diagonal = storage.copy()
    diagonal[:-1] += conductances
    diagonal[1:] += conductances
    rhs = storage * field_previous

    surface_area_per_length = 2.0 * math.pi * RADIUS_M
    boundary_conductance = surface_area_per_length / (
        (RADIUS_M - grid.radii_m[-1]) / diffusion_coefficient[-1]
        + 1.0 / boundary_transfer_coefficient
    )
    diagonal[-1] += boundary_conductance
    rhs[-1] += boundary_conductance * boundary_value
    return lower, diagonal, upper, rhs, boundary_conductance


def reconstruct_surface_value(
    cell_value: float,
    boundary_value: float,
    cell_coefficient: float,
    transfer_coefficient: float,
    cell_distance_m: float,
) -> float:
    cell_side = cell_coefficient / cell_distance_m
    return (
        cell_side * cell_value + transfer_coefficient * boundary_value
    ) / (cell_side + transfer_coefficient)


def simulate_coupled(
    environment: EnvironmentData,
    nominal_radial_step_cm: float,
    time_step_s: float,
    picard_temperature_tolerance_c: float = 1.0e-8,
    picard_moisture_tolerance: float = 1.0e-10,
    maximum_picard_iterations: int = 30,
) -> CoupledResult:
    """Solve the two-way property-coupled heat and moisture equations."""
    steps = int(round(END_TIME_S / time_step_s))
    if not math.isclose(steps * time_step_s, END_TIME_S, abs_tol=1e-12):
        raise ValueError("The time step must divide 10800 s exactly.")
    record_stride = int(round(1.0 / time_step_s))
    if not math.isclose(record_stride * time_step_s, 1.0, abs_tol=1e-12):
        raise ValueError("The time step must divide the 1 s output interval exactly.")

    grid = make_grid(nominal_radial_step_cm)
    temperature = np.full(grid.radii_m.size, INITIAL_TEMPERATURE_C, dtype=float)
    moisture = np.full(grid.radii_m.size, INITIAL_MOISTURE_KG_KG, dtype=float)
    record_times = np.arange(0.0, END_TIME_S + 1.0, 1.0)
    temperature_records = np.empty((record_times.size, temperature.size), dtype=float)
    moisture_records = np.empty_like(temperature_records)
    surface_temperature_records = np.empty(record_times.size, dtype=float)
    surface_moisture_records = np.empty(record_times.size, dtype=float)
    temperature_records[0] = temperature
    moisture_records[0] = moisture
    surface_temperature_records[0] = INITIAL_TEMPERATURE_C
    surface_moisture_records[0] = INITIAL_MOISTURE_KG_KG

    record_index = 1
    maximum_iterations_used = 0
    integrated_heat_inflow = 0.0
    integrated_moisture_outflow = 0.0
    surface_distance = RADIUS_M - grid.radii_m[-1]

    for step in range(1, steps + 1):
        time_now = step * time_step_s
        oven_temperature = float(np.interp(time_now, environment.times_s, environment.temperature_c))
        oven_moisture = float(np.interp(time_now, environment.times_s, environment.moisture_kg_kg))
        temperature_previous = temperature.copy()
        moisture_previous = moisture.copy()
        temperature_iterate = temperature_previous.copy()
        moisture_iterate = moisture_previous.copy()

        for iteration in range(1, maximum_picard_iterations + 1):
            density = density_kg_m3(moisture_iterate)
            heat_capacity = heat_capacity_j_kg_k(moisture_iterate)
            conductivity = conductivity_w_m_k(moisture_iterate)
            heat_system = build_diffusion_system(
                grid,
                temperature_previous,
                conductivity,
                density * heat_capacity,
                HEAT_TRANSFER_COEFFICIENT_W_M2_K,
                oven_temperature,
                time_step_s,
            )
            temperature_updated = solve_tridiagonal(*heat_system[:4])

            diffusivity = moisture_diffusivity_m2_s(
                moisture_iterate, temperature_updated
            )
            moisture_system = build_diffusion_system(
                grid,
                moisture_previous,
                diffusivity,
                np.ones_like(moisture_iterate),
                MASS_TRANSFER_COEFFICIENT_M_S,
                oven_moisture,
                time_step_s,
            )
            moisture_updated = solve_tridiagonal(*moisture_system[:4])

            temperature_error = float(
                np.max(np.abs(temperature_updated - temperature_iterate))
            )
            moisture_error = float(
                np.max(np.abs(moisture_updated - moisture_iterate))
            )
            temperature_iterate = temperature_updated
            moisture_iterate = moisture_updated
            if (
                temperature_error < picard_temperature_tolerance_c
                and moisture_error < picard_moisture_tolerance
            ):
                break
        else:
            raise RuntimeError(f"Coupled Picard iteration failed at t={time_now:g} s.")

        maximum_iterations_used = max(maximum_iterations_used, iteration)
        temperature = temperature_iterate
        moisture = moisture_iterate
        final_conductivity = conductivity_w_m_k(moisture)
        final_diffusivity = moisture_diffusivity_m2_s(moisture, temperature)
        surface_temperature = reconstruct_surface_value(
            temperature[-1],
            oven_temperature,
            final_conductivity[-1],
            HEAT_TRANSFER_COEFFICIENT_W_M2_K,
            surface_distance,
        )
        surface_moisture = reconstruct_surface_value(
            moisture[-1],
            oven_moisture,
            final_diffusivity[-1],
            MASS_TRANSFER_COEFFICIENT_M_S,
            surface_distance,
        )
        heat_boundary_conductance = build_diffusion_system(
            grid,
            temperature_previous,
            final_conductivity,
            density_kg_m3(moisture) * heat_capacity_j_kg_k(moisture),
            HEAT_TRANSFER_COEFFICIENT_W_M2_K,
            oven_temperature,
            time_step_s,
        )[-1]
        moisture_boundary_conductance = build_diffusion_system(
            grid,
            moisture_previous,
            final_diffusivity,
            np.ones_like(moisture),
            MASS_TRANSFER_COEFFICIENT_M_S,
            oven_moisture,
            time_step_s,
        )[-1]
        integrated_heat_inflow += (
            time_step_s
            * heat_boundary_conductance
            * (oven_temperature - temperature[-1])
        )
        integrated_moisture_outflow += (
            time_step_s
            * moisture_boundary_conductance
            * (moisture[-1] - oven_moisture)
        )

        if step % record_stride == 0:
            temperature_records[record_index] = temperature
            moisture_records[record_index] = moisture
            surface_temperature_records[record_index] = surface_temperature
            surface_moisture_records[record_index] = surface_moisture
            record_index += 1

    return CoupledResult(
        times_s=record_times,
        grid=grid,
        temperature_c=temperature_records,
        moisture_kg_kg=moisture_records,
        surface_temperature_c=surface_temperature_records,
        surface_moisture_kg_kg=surface_moisture_records,
        integrated_heat_inflow_w_s_m=integrated_heat_inflow,
        integrated_moisture_outflow_m3_kg_kg_m=integrated_moisture_outflow,
        maximum_picard_iterations=maximum_iterations_used,
    )


def sample_result(
    result: CoupledResult,
    field_name: str,
    query_times_s: np.ndarray,
    query_radii_cm: np.ndarray,
) -> np.ndarray:
    """Sample cell-centred fields with second-order centre reconstruction."""
    if field_name == "temperature":
        field = result.temperature_c
        surface = result.surface_temperature_c
    elif field_name == "moisture":
        field = result.moisture_kg_kg
        surface = result.surface_moisture_kg_kg
    else:
        raise ValueError(f"Unknown field: {field_name}")

    time_indices = np.rint(query_times_s).astype(int)
    if not np.allclose(result.times_s[time_indices], query_times_s):
        raise ValueError("Requested times are absent from the 1 s records.")
    query_radii_m = query_radii_cm / 100.0
    extended_radii = np.concatenate(([0.0], result.grid.radii_m, [RADIUS_M]))
    sampled = np.empty((query_times_s.size, query_radii_cm.size), dtype=float)
    r0_squared = result.grid.radii_m[0] ** 2
    r1_squared = result.grid.radii_m[1] ** 2
    for row_index, time_index in enumerate(time_indices):
        values = field[time_index]
        center_value = (
            values[0] * r1_squared - values[1] * r0_squared
        ) / (r1_squared - r0_squared)
        extended_values = np.concatenate(
            ([center_value], values, [surface[time_index]])
        )
        sampled[row_index] = np.interp(
            query_radii_m, extended_radii, extended_values
        )
    return sampled


def write_environment_csv(path: Path, environment: EnvironmentData) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            ["time_s", "oven_temperature_C", "oven_temperature_K", "oven_moisture_kg_kg"]
        )
        for row in zip(
            environment.times_s,
            environment.temperature_c,
            environment.temperature_k,
            environment.moisture_kg_kg,
        ):
            writer.writerow(
                [f"{row[0]:.0f}", f"{row[1]:.6f}", f"{row[2]:.6f}", f"{row[3]:.8f}"]
            )


def write_field_csv(
    path: Path, times_s: np.ndarray, radii_cm: np.ndarray, values: np.ndarray
) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", *[f"r_{radius:.1f}_cm" for radius in radii_cm]])
        for time_s, row in zip(times_s, values):
            writer.writerow([f"{time_s:.0f}", *[f"{value:.4f}" for value in row]])


def write_table_csv(
    path: Path,
    times_s: np.ndarray,
    radii_cm: np.ndarray,
    values: np.ndarray,
) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间/h", *[f"{radius:g} cm" for radius in radii_cm]])
        for time_s, row in zip(times_s, values):
            writer.writerow([f"{time_s / 3600:.1f}", *[f"{value:.4f}" for value in row]])


def build_validation(
    environment: EnvironmentData,
    coarse: CoupledResult,
    production: CoupledResult,
) -> dict:
    production_temperature = sample_result(
        production, "temperature", production.times_s, OUTPUT_RADII_CM
    )
    production_moisture = sample_result(
        production, "moisture", production.times_s, OUTPUT_RADII_CM
    )
    coarse_temperature = sample_result(
        coarse, "temperature", coarse.times_s, OUTPUT_RADII_CM
    )
    coarse_moisture = sample_result(
        coarse, "moisture", coarse.times_s, OUTPUT_RADII_CM
    )
    table_temperature = sample_result(
        production, "temperature", TABLE_TIMES_S, TABLE_RADII_CM
    )
    table_moisture = sample_result(
        production, "moisture", TABLE_TIMES_S, TABLE_RADII_CM
    )
    coarse_table_temperature = sample_result(
        coarse, "temperature", TABLE_TIMES_S, TABLE_RADII_CM
    )
    coarse_table_moisture = sample_result(
        coarse, "moisture", TABLE_TIMES_S, TABLE_RADII_CM
    )

    stored_moisture_change = float(
        np.sum(
            production.grid.volumes_m3_m
            * (production.moisture_kg_kg[-1] - production.moisture_kg_kg[0])
        )
    )
    moisture_balance_residual = (
        stored_moisture_change
        + production.integrated_moisture_outflow_m3_kg_kg_m
    )
    moisture_balance_scale = max(
        abs(stored_moisture_change),
        abs(production.integrated_moisture_outflow_m3_kg_kg_m),
        1.0e-30,
    )

    all_temperature = np.concatenate(
        (production.temperature_c.ravel(), production.surface_temperature_c)
    )
    all_moisture = np.concatenate(
        (production.moisture_kg_kg.ravel(), production.surface_moisture_kg_kg)
    )
    diffusivity = moisture_diffusivity_m2_s(
        production.moisture_kg_kg, production.temperature_c
    )
    alpha = conductivity_w_m_k(production.moisture_kg_kg) / (
        density_kg_m3(production.moisture_kg_kg)
        * heat_capacity_j_kg_k(production.moisture_kg_kg)
    )
    radial_temperature_increments = np.diff(production_temperature, axis=1)
    # Attachment 1 contains small minute-scale temperature fluctuations near
    # the plateau.  These can briefly make the outermost layer a few
    # thousandths of a degree cooler than its neighbour, so 0.01 C is used as
    # the physically meaningful radial-order tolerance.
    radial_temperature_violations = int(
        np.count_nonzero(radial_temperature_increments < -1.0e-2)
    )
    radial_moisture_violations = int(
        np.count_nonzero(np.diff(production.moisture_kg_kg, axis=1) > 1.0e-8)
    )
    time_moisture_violations = int(
        np.count_nonzero(np.diff(production.moisture_kg_kg, axis=0) > 1.0e-8)
    )

    return {
        "model_scope": {
            "time_interval_s": [0, 10800],
            "geometry": "fixed-radius one-dimensional axisymmetric cylinder",
            "latent_heat": "excluded as requested",
            "shrinkage": "excluded in question 2",
            "constitutive_laws": "Appendix 3 used from t=0 over the whole process",
            "environment": "Attachment 1 piecewise-linear interpolation at each step",
            "production_nominal_radial_step_cm": 0.0125,
            "production_time_step_s": 0.5,
            "coarse_nominal_radial_step_cm": 0.025,
            "coarse_time_step_s": 1.0,
        },
        "grid_refinement": {
            "temperature_full_max_change_C": float(
                np.max(np.abs(production_temperature - coarse_temperature))
            ),
            "temperature_full_mean_change_C": float(
                np.mean(np.abs(production_temperature - coarse_temperature))
            ),
            "moisture_full_max_change_kg_kg": float(
                np.max(np.abs(production_moisture - coarse_moisture))
            ),
            "moisture_full_mean_change_kg_kg": float(
                np.mean(np.abs(production_moisture - coarse_moisture))
            ),
            "table3_max_change_C": float(
                np.max(np.abs(table_temperature - coarse_table_temperature))
            ),
            "table4_max_change_kg_kg": float(
                np.max(np.abs(table_moisture - coarse_table_moisture))
            ),
        },
        "moisture_balance_per_unit_length": {
            "stored_change_integral": stored_moisture_change,
            "integrated_boundary_outflow": production.integrated_moisture_outflow_m3_kg_kg_m,
            "residual": moisture_balance_residual,
            "relative_residual": abs(moisture_balance_residual) / moisture_balance_scale,
        },
        "physical_ranges": {
            "temperature_min_C": float(np.min(all_temperature)),
            "temperature_max_C": float(np.max(all_temperature)),
            "moisture_min_kg_kg": float(np.min(all_moisture)),
            "moisture_max_kg_kg": float(np.max(all_moisture)),
            "diffusivity_min_m2_s": float(np.min(diffusivity)),
            "diffusivity_max_m2_s": float(np.max(diffusivity)),
            "thermal_diffusivity_min_m2_s": float(np.min(alpha)),
            "thermal_diffusivity_max_m2_s": float(np.max(alpha)),
        },
        "checks": {
            "radial_temperature_order_violation_count": radial_temperature_violations,
            "minimum_outward_temperature_increment_C": float(
                np.min(radial_temperature_increments)
            ),
            "radial_moisture_order_violation_count": radial_moisture_violations,
            "time_moisture_monotonicity_violation_count": time_moisture_violations,
            "maximum_picard_iterations": production.maximum_picard_iterations,
        },
        "selected_3h_outputs": {
            "oven_temperature_C": float(environment.temperature_c[10800]),
            "oven_moisture_kg_kg": float(environment.moisture_kg_kg[10800]),
            "center_temperature_C": float(table_temperature[-1, 0]),
            "surface_temperature_C": float(table_temperature[-1, -1]),
            "center_moisture_kg_kg": float(table_moisture[-1, 0]),
            "surface_moisture_kg_kg": float(table_moisture[-1, -1]),
            "volume_average_moisture_kg_kg": float(
                np.sum(
                    production.grid.volumes_m3_m
                    * production.moisture_kg_kg[-1]
                )
                / np.sum(production.grid.volumes_m3_m)
            ),
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, help="Path to Attachment 1 xlsx")
    parser.add_argument("--repo-root", type=Path, help="Repository root")
    args = parser.parse_args()

    script_path = Path(__file__).resolve()
    repo_root = args.repo_root.resolve() if args.repo_root else script_path.parents[1]
    data_path = args.data.resolve() if args.data else find_default_data_path(repo_root)
    result_dir = repo_root / "results" / "A_problem2_coupled"
    result_dir.mkdir(parents=True, exist_ok=True)

    environment = load_and_preprocess_environment(data_path)
    write_environment_csv(result_dir / "environment_0_3h_1s_kelvin.csv", environment)

    coarse = simulate_coupled(
        environment,
        nominal_radial_step_cm=0.025,
        time_step_s=1.0,
    )
    production = simulate_coupled(
        environment,
        nominal_radial_step_cm=0.0125,
        time_step_s=0.5,
    )

    temperature_full = sample_result(
        production, "temperature", production.times_s, OUTPUT_RADII_CM
    )
    moisture_full = sample_result(
        production, "moisture", production.times_s, OUTPUT_RADII_CM
    )
    temperature_table = sample_result(
        production, "temperature", TABLE_TIMES_S, TABLE_RADII_CM
    )
    moisture_table = sample_result(
        production, "moisture", TABLE_TIMES_S, TABLE_RADII_CM
    )

    write_field_csv(
        result_dir / "temperature_full_1s_0p1cm.csv",
        production.times_s,
        OUTPUT_RADII_CM,
        temperature_full,
    )
    write_field_csv(
        result_dir / "moisture_full_1s_0p1cm.csv",
        production.times_s,
        OUTPUT_RADII_CM,
        moisture_full,
    )
    write_table_csv(
        result_dir / "table3_temperature.csv",
        TABLE_TIMES_S,
        TABLE_RADII_CM,
        temperature_table,
    )
    write_table_csv(
        result_dir / "table4_moisture.csv",
        TABLE_TIMES_S,
        TABLE_RADII_CM,
        moisture_table,
    )
    validation = build_validation(environment, coarse, production)
    with (result_dir / "validation_summary.json").open("w", encoding="utf-8") as stream:
        json.dump(validation, stream, ensure_ascii=False, indent=2)

    print("Table 3: temperature (deg C)")
    print(np.array2string(temperature_table, precision=4, suppress_small=False))
    print("\nTable 4: moisture concentration (kg/kg)")
    print(np.array2string(moisture_table, precision=4, suppress_small=False))
    print("\nValidation summary")
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    print(f"\nResults: {result_dir}")


if __name__ == "__main__":
    main()
