"""A题问题1：附件1数据预处理与1 s分段线性插值。

本脚本只负责读取附件1、截取0--1800 s，并输出：
1. 原始观测节点；
2. 以1 s为间隔的烘房温度和水分浓度分段线性插值结果。

后续有限体积模型只读取这里生成的CSV，不再直接读取Excel。
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


END_TIME_S = 1800.0


def default_project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def default_attachment_path(project_root: Path) -> Path:
    return (
        project_root
        / "比赛题目"
        / "CUMCM2026Problems"
        / "A题"
        / "附件"
        / "附件1.xlsx"
    )


def load_attachment(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """读取附件1中第一问所需的时间、温度和水分浓度。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = [
        row[:3]
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None and float(row[0]) <= END_TIME_S
    ]
    workbook.close()

    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] != 3:
        raise ValueError("附件1未提供有效的时间、温度和水分三列数据。")
    times_s, temperature_c, moisture_kg_kg = data.T
    if times_s[0] != 0.0 or times_s[-1] != END_TIME_S:
        raise ValueError("附件1没有完整覆盖0--1800 s。")
    if not np.all(np.diff(times_s) > 0.0):
        raise ValueError("附件1时间必须严格递增。")
    if not (
        np.all(np.isfinite(temperature_c))
        and np.all(np.isfinite(moisture_kg_kg))
    ):
        raise ValueError("附件1包含缺失值或非有限数值。")
    return times_s, temperature_c, moisture_kg_kg


def piecewise_linear_interpolation(
    source_times_s: np.ndarray,
    source_values: np.ndarray,
    query_times_s: np.ndarray,
) -> np.ndarray:
    """保持原模型的np.interp分段线性插值方法。"""
    return np.interp(query_times_s, source_times_s, source_values)


def write_environment_csv(
    path: Path,
    times_s: np.ndarray,
    temperature_c: np.ndarray,
    moisture_kg_kg: np.ndarray,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "temperature_c", "moisture_kg_kg"])
        for time_s, temperature, moisture in zip(
            times_s, temperature_c, moisture_kg_kg
        ):
            # 17位有效数字可完整往返IEEE-754双精度数，避免CSV舍入改变模型结果。
            writer.writerow(
                [
                    format(float(time_s), ".17g"),
                    format(float(temperature), ".17g"),
                    format(float(moisture), ".17g"),
                ]
            )


def parse_args() -> argparse.Namespace:
    project_root = default_project_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data",
        type=Path,
        default=default_attachment_path(project_root),
        help="附件1.xlsx路径。",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=project_root / "results" / "A_problem1_modular" / "input",
        help="预处理数据输出目录。",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source_times, source_temperature, source_moisture = load_attachment(
        args.data.resolve()
    )
    query_times = np.arange(0.0, END_TIME_S + 1.0, 1.0)
    temperature_linear = piecewise_linear_interpolation(
        source_times, source_temperature, query_times
    )
    moisture_linear = piecewise_linear_interpolation(
        source_times, source_moisture, query_times
    )

    output_dir = args.output_dir.resolve()
    write_environment_csv(
        output_dir / "attachment1_observations.csv",
        source_times,
        source_temperature,
        source_moisture,
    )
    write_environment_csv(
        output_dir / "attachment1_linear_1s.csv",
        query_times,
        temperature_linear,
        moisture_linear,
    )

    metadata = {
        "method": "piecewise linear interpolation (numpy.interp)",
        "time_range_s": [0.0, END_TIME_S],
        "source_observation_count": int(source_times.size),
        "output_record_count": int(query_times.size),
        "output_interval_s": 1.0,
        "temperature_unit": "degC",
        "moisture_unit": "kg/kg",
    }
    (output_dir / "preprocessing_summary.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(metadata, ensure_ascii=False, indent=2))
    print(f"输出目录：{output_dir}")


if __name__ == "__main__":
    main()
